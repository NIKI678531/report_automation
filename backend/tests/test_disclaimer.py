"""The fixed legal disclaimer is present, immutable and traceable in every output format."""

from __future__ import annotations

import io
import re
from html import unescape

import pypdfium2 as pdfium
import pytest
from docx import Document
from sqlalchemy import select

from app.api.routes import render as render_routes, reports as report_routes
from app.domain.document import checksum
from app.domain.models import AuditEvent, Report, ReportDocument
from app.domain.page_one_presentation import PageOneLayoutOverflowError
from app.rendering import artifacts as artifact_module
from app.rendering import disclaimer as disclaimer_module
from app.rendering.disclaimer import (
    DISCLAIMER_CHECKSUM,
    DISCLAIMER_VERSION,
    DisclaimerResourceError,
    load_disclaimer,
)
from conftest import download_report


EXPECTED_DISCLAIMER_PARAGRAPHS = (
    "This document is not for public distribution outside Hong Kong. This material and the "
    "information contained in it are for professional investors as defined under the Securities "
    "and Futures Ordinance (Cap. 571 of Laws of Hong Kong) only. The investment product(s) "
    "mentioned in this material is/are authorized by the Securities and Futures Commission "
    '("SFC") in Hong Kong. Such authorization does not imply any official recommendation by the SFC.',
    "The document is intended for general information purposes only and is provided upon your "
    "request. It does not constitute any investment advice, advertisement or promotion of any "
    "investment products or any services, nor should it be construed as an offer, solicitation of "
    "offer, invitation, or recommendation to buy or sell any securities, funds, or any other "
    "financial instruments or enter into any transaction.",
    "CSOP Asset Management Limited (“CSOP”) believes that information in this document is based "
    "upon sources that are believed to be accurate, complete and reliable. However, CSOP and any "
    "of its affiliates shall not be liable for any loss, damage or expense incurred directly or "
    "indirectly by any recipient and/or its controlling shareholder as a result of the improper use "
    "of and/or reliance on this information.",
    "Investment involves risks. Past performance information presented is not indicative of future "
    "performance. Investors should refer to the Prospectus and the Product Key Facts Statement for "
    "further details, including product features and the full list of risk factors. Investors should "
    "not rely solely on this document when making investment decisions. If you wish to receive "
    "advice on investment, please consult your professional advisers.",
    "This document is prepared by CSOP and has not been reviewed by the Securities and Futures "
    "Commission in Hong Kong. This material should not be reproduced or made available to others "
    "without the written consent of CSOP.",
    "For the Index Provider Disclaimer, please refer to the Product’s offering document.",
)


def _plain_text(markup: str) -> str:
    return " ".join(unescape(re.sub(r"<[^>]+>", " ", markup)).split())


def _finalized(client, *, language_mode: str = "EN") -> dict:
    created = client.post(
        "/api/v1/reports",
        json={"report_date": "2026-06-30", "language_mode": language_mode},
    )
    assert created.status_code == 201, created.text
    report = created.json()
    finalized = client.post(
        f"/api/v1/reports/{report['id']}/finalize",
        json={"version": 1},
    )
    assert finalized.status_code == 200, finalized.text
    return finalized.json()


def _docx_text(document: Document) -> str:
    return "\n".join(
        str(node.text)
        for node in document.element.body.iter()
        if node.tag.endswith("}t") and node.text
    )


def test_approved_disclaimer_resource_is_exactly_versioned() -> None:
    resource = load_disclaimer()

    assert resource.version == DISCLAIMER_VERSION == "csop-hk-professional-investor-v1"
    assert resource.checksum == DISCLAIMER_CHECKSUM
    assert resource.title == "Disclaimer"
    assert resource.paragraphs == EXPECTED_DISCLAIMER_PARAGRAPHS
    assert resource.issuer == "Issuer: CSOP Asset Management Limited"


@pytest.mark.parametrize("language_mode", ["EN", "ZH_HANS", "ZH_HANT"])
def test_paged_preview_has_one_english_disclaimer_page_with_standard_chrome(
    client,
    language_mode: str,
) -> None:
    created = client.post(
        "/api/v1/reports",
        json={"report_date": "2026-06-30", "language_mode": language_mode},
    )
    assert created.status_code == 201, created.text

    preview = client.post(f"/api/v1/reports/{created.json()['id']}/preview")

    assert preview.status_code == 200, preview.text
    assert preview.text.count('class="report-page') == 5
    assert preview.text.count('<header class="page-header">') == 5
    assert preview.text.count('<footer class="page-footer">') == 5
    assert 'data-page="5" data-section-key="disclaimer"' in preview.text
    page_five = preview.text.split('data-page="5"', 1)[1]
    assert "Monthly Commentary | June 30, 2026" in _plain_text(page_five)
    assert '<span class="page-number">5</span>' in page_five
    assert page_five.count('alt="CSOP Asset Management"') == 1
    text = _plain_text(preview.text)
    page_five_text = _plain_text(page_five)
    positions = [page_five_text.index(paragraph) for paragraph in EXPECTED_DISCLAIMER_PARAGRAPHS]
    assert positions == sorted(positions)
    for paragraph in EXPECTED_DISCLAIMER_PARAGRAPHS:
        assert paragraph in text
    assert text.count(load_disclaimer().paragraphs[0]) == 1
    assert text.count(load_disclaimer().issuer) == 1


def test_all_exports_append_disclaimer_and_record_its_identity(client) -> None:
    report = _finalized(client)

    downloads = {
        format_name: download_report(client, report["id"], format_name)
        for format_name in ("html", "pdf", "docx")
    }

    html = downloads["html"].text
    assert '<header class="page-header">' not in html
    assert '<footer class="page-footer">' not in html
    assert 'data-section-key="disclaimer"' in html
    assert _plain_text(html).count(load_disclaimer().paragraphs[0]) == 1

    pdf = pdfium.PdfDocument(downloads["pdf"].content)
    try:
        assert len(pdf) == 5
        last_page_text = pdf[4].get_textpage().get_text_bounded()
        assert "Disclaimer" in last_page_text
        assert "not for public distribution outside Hong Kong" in last_page_text
        assert load_disclaimer().issuer in last_page_text
        assert "Monthly Commentary | June 30, 2026" in last_page_text
    finally:
        pdf.close()

    word = Document(io.BytesIO(downloads["docx"].content))
    assert len(word.sections) == 5
    word_text = _docx_text(word)
    assert word_text.count(load_disclaimer().paragraphs[0]) == 1
    assert word_text.count(load_disclaimer().issuer) == 1

    with client.app.state.testing_sessionmaker() as db:
        stored_document = db.scalar(
            select(ReportDocument).where(
                ReportDocument.report_id == report["id"],
                ReportDocument.version == report["finalized_document_version"],
            )
        )
        events = list(
            db.scalars(
                select(AuditEvent).where(
                    AuditEvent.entity_id == report["id"],
                    AuditEvent.action == "export.generated",
                )
            )
        )
    assert len(events) == 3
    assert len({event.details["content_manifest"]["checksum"] for event in events}) == 1
    for event in events:
        assert event.details["document_checksum"] == stored_document.checksum
        assert event.details["disclaimer_version"] == DISCLAIMER_VERSION
        assert event.details["disclaimer_checksum"] == DISCLAIMER_CHECKSUM
        manifest = event.details["content_manifest"]
        assert manifest["document_checksum"] == stored_document.checksum
        assert manifest["disclaimer_version"] == DISCLAIMER_VERSION
        assert manifest["disclaimer_checksum"] == DISCLAIMER_CHECKSUM
        assert manifest["checksum"] != stored_document.checksum


def test_archived_legacy_finalized_document_gets_disclaimer_and_cannot_override_it(client) -> None:
    created = client.post("/api/v1/reports", json={"report_date": "2026-06-30"}).json()
    with client.app.state.testing_sessionmaker() as db:
        report = db.get(Report, created["id"])
        document = db.scalar(
            select(ReportDocument).where(ReportDocument.report_id == created["id"])
        )
        document.template_version = "3033-v1"
        document.content = {
            **document.content,
            "template_version": "3033-v1",
            "disclaimer": {
                "title": "FORGED LEGAL COPY",
                "paragraphs": ["Do not render this."],
            },
        }
        document.checksum = checksum(document.content)
        report.template_version = "3033-v1"
        db.commit()

    finalized = client.post(
        f"/api/v1/reports/{created['id']}/finalize",
        json={"version": 1},
    )
    assert finalized.status_code == 200, finalized.text
    archived = client.delete(
        f"/api/v1/reports/{created['id']}?version={finalized.json()['version']}"
    )
    assert archived.status_code == 204, archived.text

    downloads = {
        format_name: download_report(client, created["id"], format_name)
        for format_name in ("html", "pdf", "docx")
    }
    html = downloads["html"].text

    assert "FORGED LEGAL COPY" not in html
    assert "Do not render this." not in html
    assert _plain_text(html).count(load_disclaimer().paragraphs[0]) == 1

    pdf = pdfium.PdfDocument(downloads["pdf"].content)
    try:
        assert len(pdf) == 5
        pdf_text = "\n".join(page.get_textpage().get_text_bounded() for page in pdf)
        assert "FORGED LEGAL COPY" not in pdf_text
        assert "Do not render this." not in pdf_text
        assert _plain_text(pdf_text).count(_plain_text(load_disclaimer().paragraphs[0])) == 1
    finally:
        pdf.close()

    word = Document(io.BytesIO(downloads["docx"].content))
    assert len(word.sections) == 5
    word_text = _docx_text(word)
    assert "FORGED LEGAL COPY" not in word_text
    assert "Do not render this." not in word_text
    assert word_text.count(load_disclaimer().paragraphs[0]) == 1

    with client.app.state.testing_sessionmaker() as db:
        stored_document = db.scalar(
            select(ReportDocument).where(
                ReportDocument.report_id == created["id"],
                ReportDocument.version == finalized.json()["finalized_document_version"],
            )
        )
        events = list(
            db.scalars(
                select(AuditEvent).where(
                    AuditEvent.entity_id == created["id"],
                    AuditEvent.action == "export.generated",
                )
            )
        )
    assert len(events) == 3
    assert stored_document is not None
    assert stored_document.template_version == "3033-v1"
    for event in events:
        manifest = event.details["content_manifest"]
        assert manifest["document_checksum"] == stored_document.checksum
        assert manifest["disclaimer_version"] == DISCLAIMER_VERSION
        assert manifest["disclaimer_checksum"] == DISCLAIMER_CHECKSUM
        assert manifest["checksum"] != stored_document.checksum


def test_invalid_disclaimer_resource_blocks_export_with_specific_audit_code(client, monkeypatch) -> None:
    report = _finalized(client)

    def fail(*_args, **_kwargs):
        raise DisclaimerResourceError("private deployment path")

    monkeypatch.setattr(render_routes, "generate_export", fail)
    grant = client.get(f"/api/v1/reports/{report['id']}/exports/html/download").json()
    response = client.get(grant["download_url"])

    assert response.status_code == 503
    assert response.json()["error_code"] == "DISCLAIMER_RESOURCE_INVALID"
    assert "private deployment path" not in response.text
    with client.app.state.testing_sessionmaker() as db:
        event = db.scalar(
            select(AuditEvent).where(
                AuditEvent.entity_id == report["id"],
                AuditEvent.action == "export.failed",
            )
        )
    assert event.details["error_code"] == "DISCLAIMER_RESOURCE_INVALID"
    assert event.details["disclaimer_version"] == DISCLAIMER_VERSION
    assert event.details["disclaimer_checksum"] == DISCLAIMER_CHECKSUM


def test_invalid_disclaimer_resource_blocks_preview_with_specific_code(client, monkeypatch) -> None:
    report = client.post("/api/v1/reports", json={"report_date": "2026-06-30"}).json()

    def fail(*_args, **_kwargs):
        raise DisclaimerResourceError("private deployment path")

    monkeypatch.setattr(report_routes, "render_html", fail)
    response = client.post(f"/api/v1/reports/{report['id']}/preview")

    assert response.status_code == 503
    assert response.json()["error_code"] == "DISCLAIMER_RESOURCE_INVALID"
    assert "private deployment path" not in response.text


def test_invalid_disclaimer_resource_blocks_finalize_with_specific_code(client, monkeypatch) -> None:
    report = client.post("/api/v1/reports", json={"report_date": "2026-06-30"}).json()

    def fail(*_args, **_kwargs):
        raise DisclaimerResourceError("private deployment path")

    monkeypatch.setattr(artifact_module, "assert_page_one_layout_fits", fail)
    response = client.post(
        f"/api/v1/reports/{report['id']}/finalize",
        json={"version": 1},
    )

    assert response.status_code == 503
    assert response.json()["error_code"] == "DISCLAIMER_RESOURCE_INVALID"
    assert "private deployment path" not in response.text


def test_missing_disclaimer_file_is_rejected(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(
        disclaimer_module,
        "DISCLAIMER_RESOURCE_PATH",
        tmp_path / "missing-disclaimer.json",
    )
    disclaimer_module.load_disclaimer.cache_clear()
    try:
        with pytest.raises(DisclaimerResourceError):
            disclaimer_module.load_disclaimer()
    finally:
        disclaimer_module.load_disclaimer.cache_clear()


def test_page_one_overflow_blocks_every_export_format(client, monkeypatch) -> None:
    report = _finalized(client)

    def reject_layout(*_args, **_kwargs):
        raise PageOneLayoutOverflowError(({
            "layout_id": "review:summary",
            "bottom": 812,
            "safe_bottom": 790,
        },))

    monkeypatch.setattr(artifact_module, "assert_page_one_layout_fits", reject_layout)
    for format_name in ("html", "pdf", "docx"):
        grant = client.get(
            f"/api/v1/reports/{report['id']}/exports/{format_name}/download"
        ).json()
        response = client.get(grant["download_url"])
        assert response.status_code == 422
        assert response.json()["error_code"] == "PAGE_ONE_LAYOUT_OVERFLOW"
        assert response.json()["findings"][0]["layout_id"] == "review:summary"
