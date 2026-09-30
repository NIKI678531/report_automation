from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pypdfium2 as pdfium
import pytest
from docx import Document
from sqlalchemy import select

from app.domain.document import initial_document
from app.domain.models import ProductCatalog, Report, ReportStatus
from app.domain.page_one_presentation import PageOneLayoutOverflowError
from app.rendering import artifacts as artifact_module
from app.rendering.html import render_html
from conftest import download_report


def _layout_markup(*, node_bottom_px: float) -> str:
    return f"""<!doctype html><style>
      .report-page {{ position:relative; width:400px; height:600px; overflow:hidden }}
      .page-body {{ position:relative; width:300px; height:550px }}
      .page-footer {{ position:absolute; left:0; right:0; top:550px; height:50px }}
      [data-layout-id] {{ position:absolute; left:0; top:500px; width:200px;
        height:{node_bottom_px - 500}px }}
    </style><section class="report-page" data-page="2" data-section-key="month_in_review">
      <main class="page-body"><div data-layout-id="historical_performance">History</div></main>
      <footer class="page-footer">Footer</footer>
    </section>"""


def _horizontal_layout_markup(*, clipped: bool) -> str:
    child = (
        '<div style="width:340px;height:20px">clipped table content</div>'
        if clipped
        else '<svg width="200" height="60"><title>Sector weights</title><text class="donut-label-outside" '
        'x="8" y="24">Information Technology</text></svg>'
    )
    overflow = "hidden" if clipped else "visible"
    return f"""<!doctype html><style>
      .report-page {{ position:relative; width:400px; height:600px; overflow:hidden }}
      .page-body {{ position:relative; width:300px; height:500px }}
      .page-footer {{ position:absolute; left:0; right:0; top:550px; height:50px }}
      [data-layout-id] {{ width:300px; height:60px; overflow-x:{overflow} }}
    </style><section class="report-page" data-page="4">
      <main class="page-body"><div data-layout-id="analytics">{child}</div></main>
      <footer class="page-footer">Footer</footer>
    </section>"""


def _v4_document_with_history_on(page: int) -> tuple[SimpleNamespace, dict]:
    content = initial_document(
        "report-v4",
        date(2026, 8, 31),
        "3033-v4",
        "3033-v4",
        "3033.HK",
        "HSTECHN Index",
        "EN",
    )
    content["presentation"]["page_one"]["review_page_count"] = page
    for element in content["presentation"]["page_one"]["elements"]:
        if element["id"] in {"historical_performance", "footnote:historical"}:
            element["page"] = page
    report = SimpleNamespace(
        report_date=date(2026, 8, 31),
        language_mode="EN",
        product_name="CSOP Hang Seng TECH Index ETF",
        product_code="3033",
        benchmark_code="HSTECHN",
    )
    return report, content


def _legacy_v3_document() -> tuple[SimpleNamespace, dict]:
    report, _ = _v4_document_with_history_on(1)
    content = initial_document(
        "report-v3",
        date(2026, 8, 31),
        "3033-v3",
        "3033-v3",
        "3033.HK",
        "HSTECHN Index",
        "EN",
    )
    return report, content


@pytest.mark.parametrize("bottom_px", [549.0, 550.0])
def test_vertical_overflow_up_to_three_mm_is_an_audited_warning(monkeypatch, bottom_px: float):
    monkeypatch.setattr(
        artifact_module,
        "render_html",
        lambda *_args, **_kwargs: _layout_markup(node_bottom_px=bottom_px),
    )

    findings = artifact_module.assert_page_one_layout_fits(
        None, {"template_version": "3033-v4"}
    )

    assert len(findings) == 1
    assert findings[0] == {
        "code": "PAGE_LAYOUT_TOLERANCE_USED",
        "severity": "WARNING",
        "page": 2,
        "element_id": "historical_performance",
        "layout_id": "historical_performance",
        "axis": "vertical",
        "overflow_mm": pytest.approx(findings[0]["overflow_mm"], abs=0.001),
        "message": findings[0]["message"],
    }
    assert 2.7 <= findings[0]["overflow_mm"] <= 3.0


def test_vertical_overflow_above_three_mm_remains_blocking(monkeypatch):
    monkeypatch.setattr(
        artifact_module,
        "render_html",
        lambda *_args, **_kwargs: _layout_markup(node_bottom_px=550.6),
    )

    with pytest.raises(PageOneLayoutOverflowError) as caught:
        artifact_module.assert_page_one_layout_fits(
            None, {"template_version": "3033-v4"}
        )

    finding = caught.value.findings[0]
    assert finding["code"] == "PAGE_LAYOUT_VERTICAL_OVERFLOW"
    assert finding["severity"] == "BLOCKING"
    assert finding["page"] == 2
    assert finding["overflow_mm"] > 3


@pytest.mark.parametrize(
    ("overflow_mm", "expected_severity"),
    [(0.0, None), (2.9, "WARNING"), (3.0, "WARNING"), (3.1, "BLOCKING")],
)
def test_vertical_policy_exact_boundaries(monkeypatch, overflow_mm, expected_severity):
    px_per_mm = 96 / 25.4
    soft_bottom_px = 550 - 3 * px_per_mm
    monkeypatch.setattr(
        artifact_module,
        "render_html",
        lambda *_args, **_kwargs: _layout_markup(
            node_bottom_px=soft_bottom_px + overflow_mm * px_per_mm
        ),
    )

    if expected_severity == "BLOCKING":
        with pytest.raises(PageOneLayoutOverflowError) as caught:
            artifact_module.assert_page_one_layout_fits(
                None, {"template_version": "3033-v4"}
            )
        assert caught.value.findings[0]["severity"] == expected_severity
        assert caught.value.findings[0]["overflow_mm"] == pytest.approx(overflow_mm, abs=0.005)
        return

    findings = artifact_module.assert_page_one_layout_fits(
        None, {"template_version": "3033-v4"}
    )
    if expected_severity is None:
        assert findings == ()
    else:
        assert findings[0]["severity"] == expected_severity
        assert findings[0]["overflow_mm"] == pytest.approx(overflow_mm, abs=0.005)


def test_svg_text_scroll_metrics_do_not_create_horizontal_false_positive(monkeypatch):
    monkeypatch.setattr(
        artifact_module,
        "render_html",
        lambda *_args, **_kwargs: _horizontal_layout_markup(clipped=False),
    )

    assert artifact_module.assert_page_one_layout_fits(
        None, {"template_version": "3033-v4"}
    ) == ()


def test_real_horizontal_clipping_remains_blocking(monkeypatch):
    monkeypatch.setattr(
        artifact_module,
        "render_html",
        lambda *_args, **_kwargs: _horizontal_layout_markup(clipped=True),
    )

    with pytest.raises(PageOneLayoutOverflowError) as caught:
        artifact_module.assert_page_one_layout_fits(
            None, {"template_version": "3033-v4"}
        )

    assert caught.value.findings[0]["code"] == "PAGE_LAYOUT_HORIZONTAL_OVERFLOW"
    assert caught.value.findings[0]["axis"] == "horizontal"


def test_v4_html_inserts_opening_pages_before_the_four_fixed_business_pages():
    report, content = _v4_document_with_history_on(3)

    html = render_html(report, content, layout_mode="paged")

    assert html.count('data-section-key="month_in_review"') == 3
    assert 'data-page="3" data-section-key="month_in_review" data-review-page-index="2"' in html
    assert 'data-page="4" data-section-key="company_news"' in html
    assert 'data-page="5" data-section-key="constituents"' in html
    assert 'data-page="6" data-section-key="analytics"' in html
    assert 'data-page="7" data-section-key="disclaimer"' in html
    page_three = html.index('data-page="3" data-section-key="month_in_review"')
    company_news = html.index('data-page="4" data-section-key="company_news"')
    assert page_three < html.index('data-layout-id="historical_performance"', page_three) < company_news
    assert page_three < html.index('data-layout-id="footnote:historical"', page_three) < company_news


def test_v4_continuous_html_omits_empty_opening_page_shells():
    report, content = _v4_document_with_history_on(96)

    html = render_html(report, content, layout_mode="continuous")

    assert html.count('data-section-key="month_in_review"') == 2
    assert 'data-page="1" data-section-key="month_in_review"' in html
    assert 'data-page="96" data-section-key="month_in_review"' in html
    assert 'data-page="2" data-section-key="month_in_review"' not in html


def test_v4_history_group_width_is_shared_by_html_and_docx(tmp_path):
    report, content = _v4_document_with_history_on(2)
    for element in content["presentation"]["page_one"]["elements"]:
        if element["id"] in {"historical_performance", "footnote:historical"}:
            element["x"] = 3
            element["w"] = 6

    html = render_html(report, content, layout_mode="paged")
    assert "grid-column:4 / span 6" in html

    destination = tmp_path / "v4-half-width-history.docx"
    artifact_module.render_docx(report, content, destination)
    with ZipFile(destination) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert '<w:gridSpan w:val="6"' in document_xml


def test_v4_docx_preserves_requested_inline_font_name(tmp_path):
    report, content = _v4_document_with_history_on(1)
    summary = next(
        block
        for block in content["sections"]["month_in_review"]["blocks"]
        if block["block_id"] == "summary"
    )
    summary["content"] = (
        '<p><span data-font-family="Unbundled Font">Requested face</span></p>'
    )
    destination = tmp_path / "v4-requested-font.docx"

    artifact_module.render_docx(report, content, destination)

    with ZipFile(destination) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:ascii="Unbundled Font"' in document_xml
    assert 'w:hAnsi="Unbundled Font"' in document_xml


def test_v4_history_title_keeps_reference_center_alignment_in_html_and_docx(tmp_path):
    report, content = _v4_document_with_history_on(1)

    html = render_html(report, content, layout_mode="paged")
    historical_heading = html.split('<section class="historical-module"', 1)[1].split(
        "</h2>", 1
    )[0]
    assert "text-align:center" in historical_heading

    destination = tmp_path / "v4-centered-history-title.docx"
    artifact_module.render_docx(report, content, destination)
    document = Document(destination)
    title_paragraph = next(
        paragraph
        for paragraph in document.element.body.xpath(".//w:p")
        if "Historical Performance" in "".join(paragraph.itertext())
    )
    assert title_paragraph.xpath("./w:pPr/w:jc/@w:val") == ["center"]


def test_v4_title_content_gap_is_shared_by_review_history_html_and_docx(tmp_path):
    report, content = _v4_document_with_history_on(1)
    for element in content["presentation"]["page_one"]["elements"]:
        if element["id"].startswith("review:") or element["id"] == "historical_performance":
            element["title_style"]["space_after_pt"] = 0

    html = render_html(report, content, layout_mode="paged")
    summary_heading = html.split('data-block-id="summary"', 1)[1].split("</h3>", 1)[0]
    historical_heading = html.split('<section class="historical-module"', 1)[1].split(
        "</h2>", 1
    )[0]
    assert "margin-bottom:0.0pt" in summary_heading
    assert "margin-bottom:0.0pt" in historical_heading
    assert ".history-footnote-group .history { margin-top:0; }" in html

    destination = tmp_path / "v4-zero-title-content-gap.docx"
    artifact_module.render_docx(report, content, destination)
    document = Document(destination)
    title_paragraphs = [
        paragraph
        for paragraph in document.element.body.xpath(".//w:p")
        if "August in Review" in "".join(paragraph.itertext())
        or "Historical Performance" in "".join(paragraph.itertext())
    ]
    assert len(title_paragraphs) == 2
    assert all(
        paragraph.xpath("./w:pPr/w:spacing/@w:after") == ["0"]
        for paragraph in title_paragraphs
    )


def test_v4_unbundled_inline_font_emits_explicit_pdf_fallback_warning(monkeypatch):
    report, content = _v4_document_with_history_on(1)
    summary = next(
        block
        for block in content["sections"]["month_in_review"]["blocks"]
        if block["block_id"] == "summary"
    )
    summary["content"] = (
        '<p><span data-font-family="Unbundled Font">Requested face</span></p>'
    )
    monkeypatch.setattr(
        artifact_module,
        "render_html",
        lambda *_args, **_kwargs: _horizontal_layout_markup(clipped=False),
    )

    findings = artifact_module.inspect_page_layout(report, content)

    warning = next(item for item in findings if item["code"] == "FONT_FALLBACK_USED")
    assert warning["requested_font"] == "Unbundled Font"
    assert warning["fallback_font"] == "Carlito"
    assert warning["severity"] == "WARNING"


def test_v4_docx_has_dynamic_sections_and_page_fields(tmp_path):
    report, content = _v4_document_with_history_on(3)
    destination = tmp_path / "v4-three-opening-pages.docx"

    artifact_module.render_docx(report, content, destination)

    document = Document(destination)
    assert len(document.sections) == 7
    assert all(
        len(section.footer._element.xpath(".//w:instrText")) == 1
        for section in document.sections
    )
    assert all(
        "PAGE" in section.footer._element.xpath(".//w:instrText")[0].text
        for section in document.sections
    )


def test_v3_html_keeps_the_fixed_five_page_markup():
    report, content = _legacy_v3_document()

    html = render_html(report, content, layout_mode="paged")

    assert html.count('class="report-page') == 5
    assert 'data-page="1" data-section-key="month_in_review"' in html
    assert 'data-page="5" data-section-key="disclaimer"' in html
    assert "data-review-page-index" not in html
    assert "opening-page-body" not in html


def test_v3_docx_keeps_five_sections_and_literal_page_numbers(tmp_path):
    report, content = _legacy_v3_document()
    destination = tmp_path / "legacy-v3.docx"

    artifact_module.render_docx(report, content, destination)

    document = Document(destination)
    assert len(document.sections) == 5
    assert all(
        not section.footer._element.xpath(".//w:instrText")
        for section in document.sections
    )
    for page_number, section in enumerate(document.sections, start=1):
        footer_text = "\n".join(
            cell.text
            for table in section.footer.tables
            for row in table.rows
            for cell in row.cells
        )
        assert str(page_number) in footer_text


def test_v3_preflight_keeps_page_one_only_zero_tolerance_contract(monkeypatch):
    markup = """<!doctype html><style>
      .report-page { position:relative; width:400px; height:600px; overflow:hidden }
      .page-body { position:relative; width:300px }
      .page-footer { position:absolute; top:550px }
      .footnote { position:absolute; top:540px }
      [data-layout-id] { position:absolute; top:500px; width:200px; height:49px }
    </style><section class="report-page" data-page="1">
      <main class="page-body"><div data-layout-id="review:summary">Review</div></main>
      <div class="footnote">Legacy safe boundary</div><footer class="page-footer">Footer</footer>
    </section><section class="report-page" data-page="2">
      <main class="page-body" style="height:590px">Ignored later-page overflow</main>
      <footer class="page-footer">Footer</footer>
    </section>"""
    monkeypatch.setattr(artifact_module, "render_html", lambda *_args, **_kwargs: markup)

    with pytest.raises(PageOneLayoutOverflowError) as caught:
        artifact_module.assert_page_one_layout_fits(
            None, {"template_version": "3033-v3"}
        )

    finding = caught.value.findings[0]
    assert finding["layout_id"] == "review:summary"
    assert finding["axis"] == "vertical"
    assert finding["bottom"] - finding["safe_bottom"] == pytest.approx(9)


def test_v3_pdf_keeps_the_legacy_all_page_secondary_guard(monkeypatch, tmp_path):
    markup = """<!doctype html><style>
      @page { size:A4; margin:0 }
      .report-page { position:relative; width:210mm; height:297mm; page-break-after:always }
      .page-body { height:290mm }
      .page-footer { position:absolute; top:280mm }
    </style><section class="report-page" data-page="1">
      <main class="page-body" style="height:100mm">OK</main><footer class="page-footer">Footer</footer>
    </section><section class="report-page" data-page="2">
      <main class="page-body">Overflow</main><footer class="page-footer">Footer</footer>
    </section>"""
    monkeypatch.setattr(artifact_module, "render_html", lambda *_args, **_kwargs: markup)
    document = SimpleNamespace(content={"template_version": "3033-v3"})

    with pytest.raises(ValueError, match="PDF_LAYOUT_OVERFLOW"):
        artifact_module._render_file(None, document, "pdf", tmp_path / "legacy.pdf")


def test_finalized_v3_still_exports_legacy_html_pdf_and_docx(client):
    with client.app.state.testing_sessionmaker() as db:
        product = db.scalar(
            select(ProductCatalog).where(ProductCatalog.product_code == "3033")
        )
        product.template_version = "3033-v3"
        product.design_token_version = "3033-v3"
        db.commit()
    created = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    detail = client.get(f"/api/v1/reports/{created['id']}").json()
    assert detail["latest_document"]["content"]["template_version"] == "3033-v3"
    finalized = client.post(
        f"/api/v1/reports/{created['id']}/finalize",
        json={"version": detail["latest_document"]["version"]},
    )
    assert finalized.status_code == 200, finalized.text

    html = download_report(client, created["id"], "html")
    assert html.status_code == 200
    assert html.text.count('class="report-page') == 5
    assert "data-review-page-index" not in html.text

    pdf = download_report(client, created["id"], "pdf")
    assert pdf.status_code == 200
    assert len(pdfium.PdfDocument(pdf.content)) == 5

    docx = download_report(client, created["id"], "docx")
    assert docx.status_code == 200
    word = Document(BytesIO(docx.content))
    assert len(word.sections) == 5
    assert all(
        not section.footer._element.xpath(".//w:instrText")
        for section in word.sections
    )

    with client.app.state.testing_sessionmaker() as db:
        report = db.get(Report, created["id"])
        report.status = ReportStatus.ARCHIVED
        db.commit()
    archived_html = download_report(client, created["id"], "html")
    assert archived_html.status_code == 200
    assert archived_html.content == html.content


def test_qa_blocked_can_finalize_and_layout_warning_is_audited(client, monkeypatch):
    created = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    with client.app.state.testing_sessionmaker() as db:
        report = db.get(Report, created["id"])
        report.status = ReportStatus.QA_BLOCKED
        db.commit()

    warning = {
        "code": "PAGE_LAYOUT_TOLERANCE_USED",
        "severity": "WARNING",
        "page": 1,
        "element_id": "historical_performance",
        "axis": "vertical",
        "overflow_mm": 2.9,
        "message": "Content uses 2.900 mm of the 3 mm footer tolerance.",
    }
    monkeypatch.setattr(
        artifact_module,
        "assert_page_one_layout_fits",
        lambda *_args, **_kwargs: (warning,),
    )
    document_version = client.get(
        f"/api/v1/reports/{created['id']}"
    ).json()["latest_document"]["version"]

    finalized = client.post(
        f"/api/v1/reports/{created['id']}/finalize",
        json={"version": document_version},
    )

    assert finalized.status_code == 200, finalized.text
    assert finalized.json()["status"] == "FINALIZED"
    events = client.get(f"/api/v1/audit?report_id={created['id']}").json()
    details = next(item["details"] for item in events if item["action"] == "report.finalized")
    assert details["layout_warning_findings"] == [warning]
    grant = client.get(
        f"/api/v1/reports/{created['id']}/exports/html/download"
    ).json()
    exported = client.get(grant["download_url"])
    assert exported.status_code == 200, exported.text


def test_blocking_layout_failure_is_audited_during_finalize(client, monkeypatch):
    created = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    document_version = client.get(
        f"/api/v1/reports/{created['id']}"
    ).json()["latest_document"]["version"]
    blocking = {
        "code": "PAGE_LAYOUT_VERTICAL_OVERFLOW",
        "severity": "BLOCKING",
        "page": 2,
        "element_id": "historical_performance",
        "axis": "vertical",
        "overflow_mm": 3.1,
        "message": "Content exceeds the 3 mm tolerance or touches the page footer.",
    }

    def reject(*_args, **_kwargs):
        raise PageOneLayoutOverflowError((blocking,))

    monkeypatch.setattr(artifact_module, "assert_page_one_layout_fits", reject)
    response = client.post(
        f"/api/v1/reports/{created['id']}/finalize",
        json={"version": document_version},
    )

    assert response.status_code == 422
    events = client.get(f"/api/v1/audit?report_id={created['id']}").json()
    details = next(item["details"] for item in events if item["action"] == "report.finalize_failed")
    assert details["layout_findings"] == [blocking]
    assert details["document_checksum"]
    assert details["layout_policy_version"]


def test_blocking_layout_failure_is_audited_during_export(client, monkeypatch):
    from app.api.routes import render as render_route

    created = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    document_version = client.get(
        f"/api/v1/reports/{created['id']}"
    ).json()["latest_document"]["version"]
    finalized = client.post(
        f"/api/v1/reports/{created['id']}/finalize",
        json={"version": document_version},
    )
    assert finalized.status_code == 200, finalized.text
    blocking = {
        "code": "PAGE_LAYOUT_HORIZONTAL_OVERFLOW",
        "severity": "BLOCKING",
        "page": 1,
        "element_id": "review:summary",
        "axis": "horizontal",
        "overflow_mm": 1.0,
        "message": "Content is clipped outside the horizontal page boundary.",
    }

    def reject(*_args, **_kwargs):
        raise PageOneLayoutOverflowError((blocking,))

    monkeypatch.setattr(render_route, "generate_export", reject)
    grant = client.get(
        f"/api/v1/reports/{created['id']}/exports/pdf/download"
    ).json()
    response = client.get(grant["download_url"])

    assert response.status_code == 422
    events = client.get(f"/api/v1/audit?report_id={created['id']}").json()
    details = next(
        item["details"] for item in reversed(events)
        if item["action"] == "export.failed"
    )
    assert details["layout_findings"] == [blocking]
    assert details["document_checksum"]
    assert details["layout_policy_version"]


def test_export_audits_the_same_layout_warning(client, monkeypatch):
    created = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    warning = {
        "code": "PAGE_LAYOUT_TOLERANCE_USED",
        "severity": "WARNING",
        "page": 1,
        "element_id": "historical_performance",
        "axis": "vertical",
        "overflow_mm": 2.9,
        "message": "Content uses 2.900 mm of the 3 mm footer tolerance.",
    }
    monkeypatch.setattr(
        artifact_module,
        "assert_page_one_layout_fits",
        lambda *_args, **_kwargs: (warning,),
    )
    document_version = client.get(
        f"/api/v1/reports/{created['id']}"
    ).json()["latest_document"]["version"]
    finalized = client.post(
        f"/api/v1/reports/{created['id']}/finalize",
        json={"version": document_version},
    )
    assert finalized.status_code == 200, finalized.text
    grant = client.get(
        f"/api/v1/reports/{created['id']}/exports/html/download"
    ).json()

    response = client.get(grant["download_url"])

    assert response.status_code == 200, response.text
    events = client.get(f"/api/v1/audit?report_id={created['id']}").json()
    details = next(
        item["details"] for item in reversed(events)
        if item["action"] == "export.generated"
    )
    assert details["layout_warning_findings"] == [warning]
