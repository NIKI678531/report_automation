from copy import deepcopy
from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.domain.document import bind_snapshot, initial_document, validate_document_content
from app.domain.page_one_presentation import (
    PageOneLayoutOverflowError,
    PageOnePresentationError,
    page_one_presentation,
)
from app.rendering.html import _render_tokens, render_html
from app.rendering import artifacts as artifact_module
from app.rendering.artifacts import _append_rich_text, _footer_table, render_docx


def _document() -> dict:
    content = initial_document(
        "report-1", date(2026, 8, 31), "3033-v2", "3033-v2",
        "3033.HK", "Hang Seng TECH Index", "EN",
    )
    content["sections"]["month_in_review"]["blocks"] = [
        {
            "block_id": "summary", "type": "rich_text", "title": "August in Review",
            "content": "<p>Approved review.</p>", "x": 0, "y": 0, "w": 12, "h": 4,
        },
        {
            "block_id": "drivers", "type": "key_drivers", "title": "Key Drivers",
            "content": "<p>Approved drivers.</p>", "x": 0, "y": 4, "w": 6, "h": 4,
        },
        {
            "block_id": "monitor", "type": "areas_to_monitor", "title": "Monitor",
            "content": "<p>Approved monitor.</p>", "x": 6, "y": 4, "w": 6, "h": 4,
        },
    ]
    content["sections"]["footnotes"]["historical"] = "First paragraph.\n\nSecond paragraph."
    return content


def _report():
    return SimpleNamespace(
        report_date=date(2026, 8, 31),
        language_mode="EN",
        product_name="CSOP Hang Seng TECH Index ETF",
        product_code="3033",
        benchmark_code="HSTECH",
    )


def test_new_v3_document_has_complete_valid_non_overlapping_topology():
    content = initial_document(
        "report-v3", date(2026, 8, 31), "3033-v3", "3033-v3",
        "3033.HK", "Hang Seng TECH Index", "EN",
    )

    assert [
        block["block_id"] for block in content["sections"]["month_in_review"]["blocks"]
    ] == ["summary", "drivers", "monitor", "outlook"]
    assert {
        element["id"] for element in content["presentation"]["page_one"]["elements"]
    } == {
        "review:summary", "review:drivers", "review:monitor", "review:outlook",
        "historical_performance", "footnote:historical",
    }
    assert validate_document_content(content)["presentation"]["schema_version"] == 1


def test_v2_adapter_builds_stable_non_overlapping_page_one_elements():
    adapted = page_one_presentation.adapt_document(_document())

    elements = adapted["presentation"]["page_one"]["elements"]
    assert [item["id"] for item in elements] == [
        "review:summary", "review:drivers", "review:monitor",
        "historical_performance", "footnote:historical",
    ]
    assert next(item for item in elements if item["id"] == "historical_performance")["row"] == 8
    historical = next(item for item in elements if item["id"] == "historical_performance")
    assert historical["table_font_size_role"] == "history-10"
    assert historical["table_line_height_role"] == "1.2"
    footnote = next(item for item in elements if item["id"] == "footnote:historical")
    assert [item["paragraph_index"] for item in footnote["paragraph_styles"]] == [0, 1]


@pytest.mark.parametrize("invalid_footnotes", ["bad", None, 0, []])
def test_v2_adapter_rejects_malformed_nested_content_with_presentation_error(invalid_footnotes):
    malformed = _document()
    malformed["sections"]["footnotes"] = invalid_footnotes

    with pytest.raises(PageOnePresentationError) as caught:
        page_one_presentation.adapt_document(malformed)

    assert caught.value.field == "sections.footnotes"


@pytest.mark.parametrize("invalid_blocks", [{}, "", 0, False, None])
def test_v2_adapter_does_not_treat_present_non_list_blocks_as_missing(invalid_blocks):
    malformed = _document()
    malformed["sections"]["month_in_review"]["blocks"] = invalid_blocks

    with pytest.raises(PageOnePresentationError) as caught:
        page_one_presentation.adapt_document(malformed)

    assert caught.value.field == "sections.month_in_review.blocks"


def test_fixed_page_one_elements_must_follow_the_review_flow():
    content = page_one_presentation.adapt_document(_document())
    for element in content["presentation"]["page_one"]["elements"]:
        if element["id"].startswith("review:"):
            element["row"] += 10
        elif element["id"] == "historical_performance":
            element["row"] = 0
        elif element["id"] == "footnote:historical":
            element["row"] = 1

    with pytest.raises(PageOnePresentationError, match="must follow all Review blocks"):
        validate_document_content(content)


def test_snapshot_footnote_replacement_retargets_paragraph_styles():
    content = page_one_presentation.adapt_document(_document())
    bound = bind_snapshot(
        content,
        {
            "historical_performance": {"rows": []},
            "constituents": [],
            "analytics": {"top10": [], "sectors": [], "top": [], "bottom": [], "portfolio": []},
            "footnotes": {"historical": "Replacement paragraph."},
        },
        lane="PRODUCTION",
    )
    footnote = next(
        element
        for element in bound["presentation"]["page_one"]["elements"]
        if element["id"] == "footnote:historical"
    )

    assert [style["paragraph_index"] for style in footnote["paragraph_styles"]] == [0]
    assert validate_document_content(bound)["sections"]["footnotes"]["historical"] == "Replacement paragraph."


def test_testing_snapshot_replaces_v3_placeholder_blocks_without_resetting_layout():
    from datetime import date

    content = initial_document(
        "report-1", date(2026, 8, 31), "3033-v3", "3033-v3",
        "3033.HK", "Hang Seng TECH Index", "EN",
    )
    summary_element = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    summary_element["vertical_nudge_steps"] = 2
    snapshot_review = {
        "title": "August in Review",
        "summary": "Approved fixture summary.",
        "drivers": [{"title": "Fixture driver", "body": "Driver detail."}],
        "monitor": [{"title": "Fixture monitor", "body": "Monitor detail."}],
        "outlook": "Approved fixture outlook.",
    }

    bound = bind_snapshot(
        content,
        {
            "month_in_review": snapshot_review,
            "historical_performance": {"rows": []},
            "constituents": [],
            "analytics": {"top10": [], "sectors": [], "top": [], "bottom": [], "portfolio": []},
            "footnotes": {},
        },
        lane="TESTING",
        include_testing_editorial=True,
    )

    review = bound["sections"]["month_in_review"]
    blocks = {block["block_id"]: block for block in review["blocks"]}
    assert review["summary"] == "Approved fixture summary."
    assert "Approved fixture summary." in blocks["summary"]["content"]
    assert "Fixture driver" in blocks["drivers"]["content"]
    assert "Fixture monitor" in blocks["monitor"]["content"]
    assert "Approved fixture outlook." in blocks["outlook"]["content"]
    assert next(
        item for item in bound["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )["vertical_nudge_steps"] == 2


def test_testing_snapshot_preserves_substantive_v3_review_blocks():
    from datetime import date

    content = initial_document(
        "report-1", date(2026, 8, 31), "3033-v3", "3033-v3",
        "3033.HK", "Hang Seng TECH Index", "EN",
    )
    summary = next(
        block for block in content["sections"]["month_in_review"]["blocks"]
        if block["block_id"] == "summary"
    )
    summary["content"] = "<p>Editor-authored summary.</p>"

    bound = bind_snapshot(
        content,
        {
            "month_in_review": {
                "title": "August in Review",
                "summary": "Fixture summary must not replace an edit.",
                "drivers": [],
                "monitor": [],
                "outlook": "Fixture outlook.",
            },
            "historical_performance": {"rows": []},
            "constituents": [],
            "analytics": {"top10": [], "sectors": [], "top": [], "bottom": [], "portfolio": []},
            "footnotes": {},
        },
        lane="TESTING",
        include_testing_editorial=True,
    )

    blocks = {
        block["block_id"]: block
        for block in bound["sections"]["month_in_review"]["blocks"]
    }
    assert blocks["summary"]["content"] == "<p>Editor-authored summary.</p>"


def test_presentation_is_authoritative_and_syncs_legacy_geometry():
    adapted = page_one_presentation.adapt_document(_document())
    summary = next(
        item for item in adapted["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    summary.update({"x": 1, "w": 11, "vertical_nudge_steps": 2})

    canonical = page_one_presentation.canonicalize_document(adapted)

    block = canonical["sections"]["month_in_review"]["blocks"][0]
    assert (block["x"], block["w"]) == (1, 11)
    assert summary["vertical_nudge_steps"] == 2


def test_presentation_rejects_overlap_and_arbitrary_style_values():
    adapted = page_one_presentation.adapt_document(_document())
    history = next(
        item for item in adapted["presentation"]["page_one"]["elements"]
        if item["id"] == "historical_performance"
    )
    history["row"] = 0
    with pytest.raises(PageOnePresentationError, match="overlap|must follow"):
        page_one_presentation.canonicalize_document(adapted)

    content = _document()
    content["template_version"] = "3033-v3"
    content["sections"]["month_in_review"]["blocks"][0]["content"] = (
        '<p data-font-size-role="12px">Rejected.</p>'
    )
    with pytest.raises(PageOnePresentationError, match="unsupported value"):
        validate_document_content(content)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("table_font_size_role", "history-12"),
        ("table_font_size_role", ""),
        ("table_line_height_role", "1.5"),
        ("table_line_height_role", None),
    ],
)
def test_historical_table_rejects_unsupported_style_roles(field: str, value: object):
    content = page_one_presentation.adapt_document(_document())
    historical = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "historical_performance"
    )
    historical[field] = value

    with pytest.raises(PageOnePresentationError) as caught:
        page_one_presentation.canonicalize_document(content)

    assert caught.value.field.endswith(field)
    assert caught.value.entity_id == "historical_performance"


@pytest.mark.parametrize("steps", [-200, -1, 0, 200])
def test_vertical_nudge_accepts_the_full_editor_range(steps: int):
    content = page_one_presentation.adapt_document(_document())
    summary = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    summary["vertical_nudge_steps"] = steps

    canonical = page_one_presentation.canonicalize_document(content)

    canonical_summary = next(
        item for item in canonical["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    assert canonical_summary["vertical_nudge_steps"] == steps


@pytest.mark.parametrize("steps", [-201, 201])
def test_vertical_nudge_rejects_values_outside_the_editor_range(steps: int):
    content = page_one_presentation.adapt_document(_document())
    summary = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    summary["vertical_nudge_steps"] = steps

    with pytest.raises(PageOnePresentationError) as caught:
        page_one_presentation.canonicalize_document(content)

    assert caught.value.field.endswith("vertical_nudge_steps")


def test_v3_tokens_publish_the_same_vertical_nudge_range_as_validation():
    review_tokens = _render_tokens("3033-v3")["review"]

    assert review_tokens["verticalNudgeMinSteps"] == -200
    assert review_tokens["verticalNudgeMaxSteps"] == 200


def test_reconcile_and_resolve_preserve_historical_table_style_roles():
    current = page_one_presentation.adapt_document(_document())
    current["template_version"] = "3033-v3"
    submitted = deepcopy(current)
    historical = next(
        item for item in submitted["presentation"]["page_one"]["elements"]
        if item["id"] == "historical_performance"
    )
    historical.update({
        "table_font_size_role": "history-11",
        "table_line_height_role": "1.4",
    })

    reconciled = page_one_presentation.reconcile_submission(submitted, current)
    canonical = page_one_presentation.canonicalize_document(reconciled)
    resolved = page_one_presentation.resolve_for_render(canonical, _render_tokens("3033-v3"))

    assert resolved["historical"]["table_font_size_role"] == "history-11"
    assert resolved["historical"]["table_font_size_pt"] == 11
    assert resolved["historical"]["table_line_height_role"] == "1.4"
    assert resolved["historical"]["table_line_height"] == 1.4


def test_html_historical_table_uses_resolved_font_size_and_line_height():
    content = page_one_presentation.adapt_document(_document())
    content.update(template_version="3033-v3", design_token_version="3033-v3")
    historical = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "historical_performance"
    )
    historical.update({
        "table_font_size_role": "history-11",
        "table_line_height_role": "1.4",
    })

    markup = render_html(_report(), content, layout_mode="paged")

    assert 'class="history" style="font-size:11.0pt;line-height:1.4"' in markup


def test_docx_historical_table_uses_exact_line_spacing_and_scaled_minimum_heights(tmp_path):
    content = page_one_presentation.adapt_document(_document())
    content.update(template_version="3033-v3", design_token_version="3033-v3")
    content["sections"]["historical_performance"]["rows"] = [
        {
            "name": "3033.HK",
            "return_1m": "-0.0425",
            "return_3m": "-0.0518",
            "return_6m": "-0.0979",
            "return_ytd": "-0.1610",
        }
    ]
    historical = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "historical_performance"
    )
    historical.update({
        "vertical_nudge_steps": -1,
        "table_font_size_role": "history-11",
        "table_line_height_role": "1.4",
    })
    summary = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    summary["vertical_nudge_steps"] = -1
    destination = tmp_path / "historical-styles.docx"

    render_docx(_report(), content, destination)

    document = Document(destination)
    table = next(
        item for item in document.tables
        if "1-month return (%)" in item.rows[0].cells[1].text
    )
    header_paragraph = table.rows[0].cells[1].paragraphs[0]
    body_paragraph = table.rows[1].cells[0].paragraphs[0]
    assert header_paragraph.runs[0].font.size.pt == 11
    assert body_paragraph.runs[0].font.size.pt == 11
    assert header_paragraph.paragraph_format.line_spacing.pt == pytest.approx(15.4)
    assert body_paragraph.paragraph_format.line_spacing.pt == pytest.approx(15.4)
    assert table.rows[0].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
    assert table.rows[1].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
    scale = (11 * 1.4) / (10 * 1.2)
    assert table.rows[0].height.cm == pytest.approx(0.78 * scale, abs=0.01)
    assert table.rows[1].height.cm == pytest.approx(0.62 * scale, abs=0.01)


def test_overflow_adapter_uses_stable_error_code():
    with pytest.raises(PageOneLayoutOverflowError) as caught:
        page_one_presentation.assert_no_overflow([{"layout_id": "review:summary", "bottom": 900}])
    assert caught.value.error_code == "PAGE_ONE_LAYOUT_OVERFLOW"


def test_render_resolution_maps_nudges_and_allowlisted_roles_once():
    content = page_one_presentation.adapt_document(_document())
    content["template_version"] = "3033-v3"
    summary = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "review:summary"
    )
    summary["vertical_nudge_steps"] = 3
    footnote = next(
        item for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == "footnote:historical"
    )
    footnote["bottom_nudge_steps"] = 2
    footnote["paragraph_styles"][1].update({
        "font_size_role": "footnote-10",
        "line_height_role": "1.4",
        "text_align": "justify",
    })

    resolved = page_one_presentation.resolve_for_render(content, _render_tokens("3033-v3"))

    rendered_summary = next(
        block for block in resolved["review_blocks"] if block["block_id"] == "summary"
    )
    assert rendered_summary["vertical_nudge_pt"] == 15
    assert resolved["footnote"]["bottom_nudge_pt"] == 8
    assert resolved["footnote"]["paragraphs"][1] == {
        "paragraph_index": 1,
        "text": "Second paragraph.",
        "font_size_role": "footnote-10",
        "font_size_pt": 10,
        "line_height_role": "1.4",
        "line_height": 1.4,
        "text_align": "justify",
    }


def test_docx_rich_text_consumes_the_same_paragraph_style_roles():
    document = Document()
    _append_rich_text(
        document,
        '<p data-font-size-role="review-10" data-line-height-role="1.4" '
        'data-text-align="justify">Styled paragraph.</p>',
        font_name="Calibri",
        size=11,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        style_roles=page_one_presentation.resolve_for_render(
            {**page_one_presentation.adapt_document(_document()), "template_version": "3033-v3"},
            _render_tokens("3033-v3"),
        )["style_roles"],
    )

    paragraph = document.paragraphs[0]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.line_spacing.pt == 14
    assert paragraph.runs[0].font.size.pt == 10


def test_docx_nested_list_item_keeps_its_allowlisted_paragraph_style_roles():
    document = Document()
    _append_rich_text(
        document,
        '<ul><li data-font-size-role="review-10" data-line-height-role="1.4" '
        'data-text-align="justify"><p>Styled list item.</p></li></ul>',
        font_name="Calibri",
        size=11,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        style_roles=page_one_presentation.resolve_for_render(
            {**page_one_presentation.adapt_document(_document()), "template_version": "3033-v3"},
            _render_tokens("3033-v3"),
        )["style_roles"],
    )

    paragraph = document.paragraphs[0]
    assert paragraph.text == "• Styled list item."
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.line_spacing.pt == 14
    assert all(run.font.size.pt == 10 for run in paragraph.runs)


def test_docx_numbered_list_and_plain_paragraph_keep_controlled_indents():
    document = Document()
    _append_rich_text(
        document,
        '<ol><li data-indent-level="1"><p><strong>Key driver</strong><br>'
        'Wrapped supporting copy.</p></li></ol>'
        '<p data-indent-level="2">Indented paragraph.</p>',
        font_name="Calibri",
        size=11,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    numbered, indented = document.paragraphs
    assert numbered.text == "1. Key driver\nWrapped supporting copy."
    assert numbered.paragraph_format.left_indent.pt == pytest.approx(34.7559, abs=0.01)
    assert numbered.paragraph_format.first_line_indent.cm == pytest.approx(-0.35, abs=0.01)
    assert indented.text == "Indented paragraph."
    assert indented.paragraph_format.left_indent.pt == pytest.approx(44, abs=0.01)


def test_docx_review_title_and_complete_body_use_independent_two_character_indents():
    document = Document()
    artifact_module._review_block(
        document,
        {
            "title": "Editable subheading",
            "rendered_content_html": "<p>Complete body.</p>",
            "title_style": {"font_size_pt": 12, "indent_level": 1},
            "body_style": {"font_size_pt": 10, "indent_level": 2},
        },
        font_name="Calibri",
        font_size=10,
        deep=artifact_module.RGBColor(34, 50, 127),
    )

    title, body = document.paragraphs
    assert title.text == "Editable subheading"
    assert title.paragraph_format.left_indent.pt == pytest.approx(24, abs=0.01)
    assert body.text == "Complete body."
    assert body.paragraph_format.left_indent.pt == pytest.approx(40, abs=0.01)


@pytest.mark.parametrize("steps", range(-3, 5))
def test_docx_footnote_nudge_keeps_every_allowed_position_distinct(steps: int):
    document = Document()
    _footer_table(
        document.sections[0],
        1,
        "Footnote",
        "Calibri",
        footnote_paragraphs=[{
            "text": "Footnote",
            "font_size_pt": 8,
            "line_height": 1.2,
            "text_align": "left",
        }],
        footnote_bottom_nudge_pt=steps * 4,
    )

    assert f'w:position w:val="{steps * 8}"' in document.sections[0].footer._element.xml


def test_chromium_preflight_detects_horizontal_scroll_overflow(monkeypatch):
    markup = """<!doctype html><style>
      .report-page { position:relative; width:400px; height:600px; overflow:hidden }
      .page-body { width:300px }
      .page-footer { position:absolute; top:550px }
      .footnote { position:absolute; top:500px }
      [data-layout-id] { width:100px; white-space:nowrap; overflow:hidden }
    </style><section class="report-page" data-page="1">
      <main class="page-body"><div data-layout-id="review:summary">ThisUnbrokenTextMustOverflowItsVeryNarrowContainer</div></main>
      <div class="footnote" data-layout-id="footnote:historical">Footnote</div>
      <footer class="page-footer">Footer</footer>
    </section>"""
    monkeypatch.setattr(artifact_module, "render_html", lambda *_args, **_kwargs: markup)

    with pytest.raises(PageOneLayoutOverflowError) as caught:
        artifact_module.assert_page_one_layout_fits(None, {})

    assert caught.value.findings[0]["layout_id"] == "review:summary"
    assert caught.value.findings[0]["axis"] == "horizontal"


def test_chromium_preflight_detects_actual_layout_node_overlap(monkeypatch):
    markup = """<!doctype html><style>
      .report-page { position:relative; width:400px; height:600px; overflow:hidden }
      .page-body { width:300px }
      .page-footer { position:absolute; top:550px }
      .footnote { position:absolute; top:500px; height:20px }
      [data-layout-id] { width:200px; height:100px }
      .historical-module { margin-top:-50px }
    </style><section class="report-page" data-page="1">
      <main class="page-body">
        <section data-layout-id="review:summary">Review</section>
        <section class="historical-module" data-layout-id="historical_performance">History</section>
      </main>
      <div class="footnote" data-layout-id="footnote:historical">Footnote</div>
      <footer class="page-footer">Footer</footer>
    </section>"""
    monkeypatch.setattr(artifact_module, "render_html", lambda *_args, **_kwargs: markup)

    with pytest.raises(PageOneLayoutOverflowError) as caught:
        artifact_module.assert_page_one_layout_fits(None, {})

    overlap = next(finding for finding in caught.value.findings if finding["axis"] == "overlap")
    assert overlap["layout_id"] == "review:summary"
    assert overlap["overlaps_with"] == "historical_performance"


def test_unsaved_preview_is_canonical_but_does_not_create_a_version_or_audit(client):
    report = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    detail = client.get(f"/api/v1/reports/{report['id']}").json()
    original = deepcopy(detail["latest_document"])
    content = detail["latest_document"]["content"]
    content["sections"]["month_in_review"]["blocks"] = [{
        "block_id": "summary", "type": "rich_text", "title": "Preview only",
        "content": '<p data-font-size-role="review-11" data-line-height-role="1.2" '
        'data-text-align="justify">Unsaved canonical preview.</p>',
        "x": 0, "y": 0, "w": 12, "h": 4,
    }]
    before_events = client.get(f"/api/v1/audit?report_id={report['id']}").json()

    response = client.post(
        f"/api/v1/reports/{report['id']}/preview",
        json={"version": original["version"], "content": content},
    )

    assert response.status_code == 200, response.text
    assert "Unsaved canonical preview." in response.text
    assert 'data-layout-id="review:summary"' in response.text
    assert 'data-layout-id="historical_performance"' in response.text
    assert 'data-layout-id="footnote:historical"' in response.text
    assert 'data-font-size-role="review-11"' in response.text
    after = client.get(f"/api/v1/reports/{report['id']}").json()["latest_document"]
    after_events = client.get(f"/api/v1/audit?report_id={report['id']}").json()
    assert after == original
    assert after_events == before_events


def test_draft_preview_is_version_tolerant_and_first_save_upgrades_to_v4_atomically(client):
    report = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    detail = client.get(f"/api/v1/reports/{report['id']}").json()
    content = detail["latest_document"]["content"]

    preview = client.post(
        f"/api/v1/reports/{report['id']}/preview",
        json={"version": 999, "content": content},
    )
    assert preview.status_code == 200, preview.text

    saved = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": detail["latest_document"]["version"], "content": content},
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["content"]["template_version"] == "3033-v4"
    assert saved.json()["content"]["presentation"]["schema_version"] == 2
    refreshed = client.get(f"/api/v1/reports/{report['id']}").json()
    assert refreshed["template_version"] == "3033-v4"


def test_v2_submission_without_presentation_reports_nested_shape_errors_as_422(client):
    report = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    detail = client.get(f"/api/v1/reports/{report['id']}").json()
    assert detail["latest_document"]["content"]["template_version"] == "3033-v2"

    malformed_payloads = []
    malformed_sections = deepcopy(detail["latest_document"]["content"])
    malformed_sections.pop("presentation", None)
    malformed_sections["sections"] = []
    malformed_payloads.append((malformed_sections, "sections"))
    malformed_blocks = deepcopy(detail["latest_document"]["content"])
    malformed_blocks.pop("presentation", None)
    malformed_blocks["sections"]["month_in_review"]["blocks"] = 1
    malformed_payloads.append((malformed_blocks, "sections.month_in_review.blocks"))
    falsey_blocks = deepcopy(detail["latest_document"]["content"])
    falsey_blocks.pop("presentation", None)
    falsey_blocks["sections"]["month_in_review"]["blocks"] = {}
    malformed_payloads.append((falsey_blocks, "sections.month_in_review.blocks"))
    malformed_footnotes = deepcopy(detail["latest_document"]["content"])
    malformed_footnotes.pop("presentation", None)
    malformed_footnotes["sections"]["footnotes"] = "bad"
    malformed_payloads.append((malformed_footnotes, "sections.footnotes"))
    null_footnotes = deepcopy(detail["latest_document"]["content"])
    null_footnotes.pop("presentation", None)
    null_footnotes["sections"]["footnotes"] = None
    malformed_payloads.append((null_footnotes, "sections.footnotes"))
    whitespace_id = deepcopy(detail["latest_document"]["content"])
    whitespace_id["sections"]["month_in_review"]["blocks"].append({
        "block_id": " new ", "type": "rich_text", "title": "New",
        "content": "<p>New.</p>", "x": 0, "y": 13, "w": 12, "h": 1,
        "text_align": "left",
    })
    malformed_payloads.append((
        whitespace_id,
        f"sections.month_in_review.blocks.{len(whitespace_id['sections']['month_in_review']['blocks']) - 1}.block_id",
    ))
    malformed_new_element = deepcopy(detail["latest_document"]["content"])
    malformed_new_element["sections"]["month_in_review"]["blocks"][0].update({"x": 1, "w": 11})
    malformed_new_element["sections"]["month_in_review"]["blocks"].append({
        "block_id": "new", "type": "rich_text", "title": "New",
        "content": "<p>New.</p>", "x": 0, "y": 13, "w": 12, "h": 1,
        "text_align": "left",
    })
    malformed_new_element["presentation"]["page_one"]["elements"].append({
        "id": "review:new", "row_span": 1, "x": 0, "w": 12,
        "vertical_nudge_steps": 0,
    })
    malformed_payloads.append((malformed_new_element, "presentation.page_one.elements"))

    for malformed_content, expected_field in malformed_payloads:
        for endpoint in ("document", "preview"):
            response = (
                client.patch(
                    f"/api/v1/reports/{report['id']}/document",
                    json={"version": detail["latest_document"]["version"], "content": malformed_content},
                )
                if endpoint == "document"
                else client.post(
                    f"/api/v1/reports/{report['id']}/preview",
                    json={"version": detail["latest_document"]["version"], "content": malformed_content},
                )
            )
            assert response.status_code == 422
            assert response.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"
            assert response.json()["field"] == expected_field


def test_v3_accepts_review_title_changes_and_rejects_invalid_presentation(client):
    report = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    detail = client.get(f"/api/v1/reports/{report['id']}").json()
    saved = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": detail["latest_document"]["version"], "content": detail["latest_document"]["content"]},
    )
    assert saved.status_code == 200, saved.text

    titled = deepcopy(saved.json()["content"])
    titled["sections"]["month_in_review"]["blocks"][0]["title"] = "Mutable title"
    titled["sections"]["month_in_review"]["title"] = "Mutable title"
    titled["sections"]["month_in_review"]["display_title"] = "Mutable title"
    retitled = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": saved.json()["version"], "content": titled},
    )
    assert retitled.status_code == 200, retitled.text
    assert retitled.json()["content"]["sections"]["month_in_review"]["blocks"][0][
        "title"
    ] == "Mutable title"
    saved = retitled

    for field, value in (
        ("product_name", "Mutable product heading"),
        ("benchmark_name", "Mutable benchmark heading"),
    ):
        renamed_heading = deepcopy(saved.json()["content"])
        renamed_heading["terminology_overrides"][field] = value
        rejected_heading = client.patch(
            f"/api/v1/reports/{report['id']}/document",
            json={"version": saved.json()["version"], "content": renamed_heading},
        )
        assert rejected_heading.status_code == 422
        assert rejected_heading.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"
        assert rejected_heading.json()["field"] == f"terminology_overrides.{field}"

    malformed_terms = deepcopy(saved.json()["content"])
    malformed_terms["terminology_overrides"] = []
    rejected_terms = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": saved.json()["version"], "content": malformed_terms},
    )
    assert rejected_terms.status_code == 422
    assert rejected_terms.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"
    assert rejected_terms.json()["field"] == "terminology_overrides"

    malformed_payloads = []
    malformed_page_one = deepcopy(saved.json()["content"])
    malformed_page_one["presentation"]["page_one"] = []
    malformed_payloads.append((malformed_page_one, "presentation.page_one"))
    for invalid_schema_version in (True, 1.0, 1):
        malformed_schema = deepcopy(saved.json()["content"])
        malformed_schema["presentation"]["schema_version"] = invalid_schema_version
        malformed_payloads.append((malformed_schema, "presentation.schema_version"))
    missing_row = deepcopy(saved.json()["content"])
    missing_row["presentation"]["page_one"]["elements"][0].pop("row")
    malformed_payloads.append((missing_row, "presentation.page_one.elements.0.row"))
    malformed_blocks = deepcopy(saved.json()["content"])
    malformed_blocks["sections"]["month_in_review"]["blocks"] = 1
    malformed_payloads.append((malformed_blocks, "sections.month_in_review.blocks"))
    malformed_footnotes = deepcopy(saved.json()["content"])
    malformed_footnotes["sections"]["footnotes"] = "bad"
    malformed_payloads.append((malformed_footnotes, "sections.footnotes"))
    for malformed_content, expected_field in malformed_payloads:
        for endpoint in ("document", "preview"):
            response = (
                client.patch(
                    f"/api/v1/reports/{report['id']}/document",
                    json={"version": saved.json()["version"], "content": malformed_content},
                )
                if endpoint == "document"
                else client.post(
                    f"/api/v1/reports/{report['id']}/preview",
                    json={"version": saved.json()["version"], "content": malformed_content},
                )
            )
            assert response.status_code == 422
            assert response.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"
            assert response.json()["field"] == expected_field

    malformed = deepcopy(saved.json()["content"])
    summary = next(
        element for element in malformed["presentation"]["page_one"]["elements"]
        if element["id"] == "review:summary"
    )
    summary["x"] = 11
    summary["w"] = 2
    rejected_geometry = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": saved.json()["version"], "content": malformed},
    )
    assert rejected_geometry.status_code == 422
    assert rejected_geometry.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"

    without_presentation = deepcopy(saved.json()["content"])
    without_presentation.pop("presentation")
    rejected_missing = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": saved.json()["version"], "content": without_presentation},
    )
    assert rejected_missing.status_code == 422
    assert rejected_missing.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"

    without_summary_element = deepcopy(saved.json()["content"])
    without_summary_element["presentation"]["page_one"]["elements"] = [
        element
        for element in without_summary_element["presentation"]["page_one"]["elements"]
        if element["id"] != "review:summary"
    ]
    rejected_element = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": saved.json()["version"], "content": without_summary_element},
    )
    assert rejected_element.status_code == 422
    assert rejected_element.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"

    renamed = deepcopy(saved.json()["content"])
    summary_block = next(
        block for block in renamed["sections"]["month_in_review"]["blocks"]
        if block["block_id"] == "summary"
    )
    summary_block.update(block_id="renamed", title="Mutable module title")
    summary_element = next(
        element for element in renamed["presentation"]["page_one"]["elements"]
        if element["id"] == "review:summary"
    )
    summary_element["id"] = "review:renamed"
    rejected_identity = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": saved.json()["version"], "content": renamed},
    )
    assert rejected_identity.status_code == 422
    assert rejected_identity.json()["error_code"] == "PAGE_ONE_PRESENTATION_INVALID"


def test_long_page_one_copy_can_be_saved_but_blocks_finalization(client):
    report = client.post("/api/v1/reports", json={"report_date": "2026-08-31"}).json()
    detail = client.get(f"/api/v1/reports/{report['id']}").json()
    content = detail["latest_document"]["content"]
    content["sections"]["month_in_review"]["blocks"] = [{
        "block_id": "summary",
        "type": "rich_text",
        "title": "August in Review",
        "content": f"<p>{'Overflow copy. ' * 800}</p>",
        "x": 0,
        "y": 0,
        "w": 12,
        "h": 4,
        "text_align": "left",
    }]

    saved = client.patch(
        f"/api/v1/reports/{report['id']}/document",
        json={"version": detail["latest_document"]["version"], "content": content},
    )
    assert saved.status_code == 200, saved.text
    refreshed = client.get(f"/api/v1/reports/{report['id']}").json()

    finalized = client.post(
        f"/api/v1/reports/{report['id']}/finalize",
        json={"version": refreshed["version"]},
    )

    assert finalized.status_code == 422, finalized.text
    assert finalized.json()["error_code"] == "PAGE_ONE_LAYOUT_OVERFLOW"
    assert any(
        finding["layout_id"] == "review:summary"
        for finding in finalized.json()["findings"]
    )
