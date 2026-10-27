from copy import deepcopy
from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.domain.document import initial_document, validate_document_content
from app.core.config import settings
from app.domain.page_one_presentation import PageOnePresentationError, page_one_presentation
from app.rendering.artifacts import render_docx
from app.rendering.html import hong_kong_export_date, render_html


@pytest.fixture
def calibri_css(monkeypatch):
    monkeypatch.setattr("app.rendering.html._embedded_calibri_font_css", lambda: "")


def _document(version: str = "3033-v5") -> dict:
    content = initial_document(
        "report-v5",
        date(2026, 6, 30),
        version,
        version,
        "3033.HK",
        "Hang Seng TECH Index",
        "EN",
    )
    content["sections"]["footnotes"]["historical"] = "Source¹: CSOP.\n\nPast performance² is not indicative."
    if content.get("presentation", {}).get("schema_version") in {2, 3}:
        _element(content, "footnote:historical")["content_html"] = (
            "<p>Source¹: CSOP.</p><p>Past performance² is not indicative.</p>"
        )
    return content


def _element(content: dict, element_id: str) -> dict:
    return next(
        item
        for item in content["presentation"]["page_one"]["elements"]
        if item["id"] == element_id
    )


def _report() -> SimpleNamespace:
    return SimpleNamespace(
        report_date=date(2026, 6, 30),
        language_mode="EN",
        product_name="CSOP Hang Seng TECH Index ETF",
        product_code="3033",
        benchmark_code="HSTECH",
    )


def test_v5_defaults_put_locked_eight_point_footnote_in_safe_footer_band():
    content = validate_document_content(_document())
    historical = _element(content, "historical_performance")
    footnote = _element(content, "footnote:historical")

    assert content["presentation"]["schema_version"] == 3
    assert (footnote["page"], footnote["x"], footnote["w"]) == (1, 0, 12)
    assert footnote["page"] == historical["page"]
    assert footnote["footer_bottom_mm"] == 20
    assert not {"row", "row_span", "gap_pt"} & footnote.keys()
    assert footnote["body_style"] == {
        "font_family": "Calibri",
        "font_size_pt": 8,
        "color": "#000000",
        "bold": False,
        "italic": False,
        "underline": False,
        "line_height": 1.2,
        "space_before_pt": 0,
        "space_after_pt": 0,
        "text_align": "left",
        "indent_level": 0,
    }


@pytest.mark.parametrize(
    ("language_mode", "text"),
    [
        ("EN", "Source¹: CSOP."),
        ("ZH_HANS", "资料来源¹：南方东英。"),
        ("ZH_HANT", "資料來源¹：南方東英。"),
    ],
)
def test_v5_preserves_multilingual_footnote_text_and_unicode_markers(
    language_mode: str,
    text: str,
):
    content = initial_document(
        f"report-{language_mode}",
        date(2026, 6, 30),
        "3033-v5",
        "3033-v5",
        "3033.HK",
        "Hang Seng TECH Index",
        language_mode,
    )
    content["sections"]["footnotes"]["historical"] = text
    _element(content, "footnote:historical")["content_html"] = f"<p>{text}</p>"

    canonical = page_one_presentation.canonicalize_document(content)

    assert canonical["sections"]["footnotes"]["historical"] == text
    assert text in render_html(_report(), canonical, layout_mode="paged")


def test_v5_locks_inline_style_but_keeps_allowed_paragraph_and_position_controls():
    content = _document()
    footnote = _element(content, "footnote:historical")
    footnote.update(x=2, w=8, footer_bottom_mm=31)
    footnote["content_html"] = (
        '<p data-line-height="2"><span data-font-size-pt="30"><strong>Source¹</strong>'
        "<br>Line 2</span></p>"
    )
    footnote["body_style"].update(
        font_family="Comic Sans MS",
        font_size_pt=30,
        color="#FF0000",
        bold=True,
        italic=True,
        underline=True,
        line_height=1.4,
        space_before_pt=2,
        space_after_pt=3,
        text_align="right",
    )

    canonical = page_one_presentation.canonicalize_document(content)
    saved = _element(canonical, "footnote:historical")

    assert (saved["x"], saved["w"], saved["footer_bottom_mm"]) == (2, 8, 31)
    assert saved["content_html"] == "<p>Source¹<br>Line 2</p>"
    assert saved["body_style"]["font_size_pt"] == 8
    assert saved["body_style"]["font_family"] == "Calibri"
    assert saved["body_style"]["color"] == "#000000"
    assert saved["body_style"]["bold"] is False
    assert saved["body_style"]["line_height"] == 1.4
    assert saved["body_style"]["space_before_pt"] == 2
    assert saved["body_style"]["space_after_pt"] == 3
    assert saved["body_style"]["text_align"] == "right"
    assert canonical["sections"]["footnotes"]["historical"] == "Source¹\nLine 2"


@pytest.mark.parametrize("footer_bottom_mm", [15.99, 36.01])
def test_v5_rejects_footer_position_outside_safe_band(footer_bottom_mm: float):
    content = _document()
    _element(content, "footnote:historical")["footer_bottom_mm"] = footer_bottom_mm

    with pytest.raises(PageOnePresentationError) as caught:
        page_one_presentation.canonicalize_document(content)

    assert caught.value.field.endswith("footer_bottom_mm")


def test_v4_draft_upgrade_preserves_page_and_columns_but_resets_vertical_position():
    v4 = _document("3033-v4")
    historical = _element(v4, "historical_performance")
    footnote = _element(v4, "footnote:historical")
    historical.update(page=2, x=3, w=7)
    footnote.update(page=2, x=3, w=7, gap_pt=19)
    footnote["content_html"] = "<p>Kept¹</p>"
    v4["presentation"]["page_one"]["review_page_count"] = 2

    upgraded = page_one_presentation.upgrade_document_to_v5(v4)
    upgraded_footnote = _element(upgraded, "footnote:historical")

    assert upgraded["template_version"] == "3033-v5"
    assert upgraded["presentation"]["schema_version"] == 3
    assert (upgraded_footnote["page"], upgraded_footnote["x"], upgraded_footnote["w"]) == (2, 3, 7)
    assert upgraded_footnote["footer_bottom_mm"] == 20
    assert upgraded_footnote["content_html"] == "<p>Kept¹</p>"
    assert "gap_pt" not in upgraded_footnote


def test_v5_paged_and_continuous_html_share_footer_style_without_body_footnote():
    content = _document()
    paged = render_html(_report(), content, layout_mode="paged")
    continuous = render_html(_report(), content, layout_mode="continuous")

    for markup in (paged, continuous):
        assert markup.count('data-layout-id="footnote:historical"') == 1
        assert 'class="review-footnote-layer"' in markup
        assert "font-size:8.0pt" in markup
        assert "Source¹" in markup


def test_v5_keeps_existing_review_typography_controls_while_locking_footnote_style():
    content = _document()
    summary = _element(content, "review:summary")
    summary["body_style"]["font_size_pt"] = 11
    summary["title_style"]["space_after_pt"] = 6

    paged = render_html(_report(), content, layout_mode="paged")

    assert "font-size:11.0pt" in paged
    assert "margin-bottom:6.0pt" in paged
    assert "font-size:8.0pt" in paged


def test_v5_historical_footnote_follows_historical_to_the_selected_review_page():
    content = _document()
    content["presentation"]["page_one"]["review_page_count"] = 2
    _element(content, "historical_performance")["page"] = 2
    _element(content, "footnote:historical")["page"] = 2

    paged = render_html(_report(), content, layout_mode="paged")
    page_one_start = paged.index('data-page="1"')
    page_two_start = paged.index('data-page="2"')
    page_three_start = paged.index('data-page="3"')
    footnote_start = paged.index('data-layout-id="footnote:historical"')

    assert page_one_start < page_two_start < footnote_start < page_three_start
    assert paged.count('data-layout-id="footnote:historical"') == 1


def test_v5_docx_writes_historical_footnote_only_to_footer_xml(tmp_path):
    content = _document()
    _element(content, "footnote:historical")["body_style"].update(
        line_height=1.4,
        space_before_pt=2,
        space_after_pt=3,
        text_align="right",
    )
    destination = tmp_path / "v5-review-footer.docx"

    render_docx(_report(), content, destination)

    with ZipFile(destination) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
    assert "Source¹" not in document_xml
    assert "Source¹" in footer_xml
    assert "Past performance²" in footer_xml

    exported = Document(destination)
    paragraph = next(
        paragraph
        for section in exported.sections
        for paragraph in section.footer.paragraphs
        if "Source¹" in paragraph.text
    )
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph.paragraph_format.space_before.pt == 2
    assert paragraph.paragraph_format.space_after.pt >= 3
    assert paragraph.paragraph_format.line_spacing.pt == pytest.approx(11.2)


def test_v5_docx_footer_follows_historical_to_its_review_section(tmp_path):
    content = _document()
    content["presentation"]["page_one"]["review_page_count"] = 2
    _element(content, "historical_performance")["page"] = 2
    _element(content, "footnote:historical")["page"] = 2
    destination = tmp_path / "v5-review-page-two.docx"

    render_docx(_report(), content, destination)

    exported = Document(destination)
    first_footer = "\n".join(paragraph.text for paragraph in exported.sections[0].footer.paragraphs)
    second_footer = "\n".join(paragraph.text for paragraph in exported.sections[1].footer.paragraphs)
    assert "Source¹" not in first_footer
    assert "Source¹" in second_footer


def test_finalized_v4_content_can_still_render_with_legacy_body_footnote():
    v4 = deepcopy(_document("3033-v4"))
    markup = render_html(_report(), v4, layout_mode="paged")

    assert 'class="review-footnote-layer"' not in markup
    assert markup.count('data-layout-id="footnote:historical"') == 1
    assert "font-size:9.0pt" in markup


@pytest.mark.parametrize("version", [f"3033-v{number}" for number in range(1, 7)])
@pytest.mark.parametrize("product_code", ["3033", "3149"])
@pytest.mark.parametrize("layout_mode", ["continuous", "paged"])
def test_reader_sidebar_is_in_all_downloaded_html(
    version, product_code, layout_mode, calibri_css, monkeypatch,
):
    from xml.etree import ElementTree

    monkeypatch.setattr("app.rendering.html.hong_kong_export_date", lambda: "26 Oct 2026")
    content = _document(version)
    report = _report()
    report.product_code = product_code

    markup = render_html(report, content, layout_mode=layout_mode)

    expected = layout_mode == "continuous"
    assert ('class="product-sidebar"' in markup) is expected
    if expected:
        sidebar = ElementTree.fromstring(
            markup[markup.index('<aside class="product-sidebar"'):markup.index('</aside>') + 8]
        )
        assert "Authored by:" in "".join(sidebar.itertext())
        assert "CSOP Intelligence" in "".join(sidebar.itertext())
        assert "Date : 26 Oct 2026" in " ".join("".join(sidebar.itertext()).split())
        assert "Product Insights" in "".join(sidebar.itertext())
        assert "Related Products" in "".join(sidebar.itertext())
        assert "Related ETFs" not in "".join(sidebar.itertext())
        assert (
            f"<h1>CSOP Hang Seng TECH Index ETF ({product_code}.HK) Monthly Commentary</h1>"
            in markup
        )
        cards = sidebar.findall(".//article")
        assert [card.find("p").text for card in cards] == ["3033", "3149"]
        links = sidebar.findall(".//a[@class='product-link']")
        assert [link.attrib["href"] for link in links] == [
            "https://www.csopasset.com/en/products/co-chst",
            "https://www.csopasset.com/en/products/hk-cptc",
        ]
        assert ["".join(link.itertext()).strip() for link in links] == [
            "View 3033 holdings, NAV & performance details >>",
            "View 3149 holdings, NAV & performance details >>",
        ]
        assert all(link.attrib["rel"] == "noopener noreferrer" for link in links)
        assert all(link.attrib["target"] == "_blank" for link in links)


def test_reader_sidebar_date_uses_the_hong_kong_export_day():
    assert hong_kong_export_date(datetime(2026, 10, 25, 16, 30, tzinfo=timezone.utc)) == "26 Oct 2026"


def test_v6_keeps_footer_geometry_and_does_not_mutate_v5_source(calibri_css):
    source = _document()
    _element(source, "footnote:historical").update(x=2, w=8, footer_bottom_mm=31)
    original = deepcopy(source)
    upgraded = page_one_presentation.upgrade_document_to_v6(source)
    canonical = validate_document_content(upgraded)
    saved = _element(canonical, "footnote:historical")

    assert source == original
    assert canonical["template_version"] == canonical["design_token_version"] == "3033-v6"
    assert canonical["presentation"]["schema_version"] == 3
    assert (saved["x"], saved["w"], saved["footer_bottom_mm"]) == (2, 8, 31)
    markup = render_html(_report(), canonical, layout_mode="paged")
    assert 'class="review-footnote-layer" style="bottom:31.0mm"' in markup
    assert markup.count('data-layout-id="footnote:historical"') == 1


def test_v6_initial_document_uses_footer_presentation():
    content = _document("3033-v6")
    assert content["presentation"]["schema_version"] == 3
    assert _element(content, "footnote:historical")["footer_bottom_mm"] == 20


def test_v6_uses_one_calibri_footnote_family(calibri_css):
    content = _document("3033-v6")
    markup = render_html(_report(), content, layout_mode="paged")
    assert '.footnote { font-family:"Report Calibri"' in markup
    assert "font-family:Report Calibri;font-size:8.0pt" in markup


def test_v6_missing_font_fails_without_exposing_filesystem_path(monkeypatch, tmp_path):
    from fastapi import HTTPException

    missing = tmp_path / "private" / "calibri.ttf"
    monkeypatch.setattr(settings, "calibri_font_path", missing)
    with pytest.raises(HTTPException) as caught:
        render_html(_report(), _document("3033-v6"))
    assert caught.value.status_code == 503
    assert caught.value.detail["error_code"] == "CALIBRI_FONT_UNAVAILABLE"
    assert str(missing) not in str(caught.value.detail)
    assert "Report Calibri" not in render_html(_report(), _document("3033-v5"))


def test_v6_docx_footer_uses_calibri_and_preserves_position(tmp_path):
    from docx.oxml.ns import qn

    content = _document("3033-v6")
    _element(content, "footnote:historical").update(x=2, w=8, footer_bottom_mm=31)
    destination = tmp_path / "v6-footer.docx"
    render_docx(_report(), content, destination)

    exported = Document(destination)
    paragraph = next(
        paragraph for paragraph in exported.sections[0].footer.paragraphs
        if "Source" in paragraph.text
    )
    assert paragraph.runs[0].font.name == "Calibri"
    assert paragraph.runs[0].font.size.pt == 8
    assert paragraph.paragraph_format.left_indent.pt > 0
    assert paragraph.runs[0]._element.find(".//" + qn("w:position")) is not None
    assert not any("Source" in paragraph.text for paragraph in exported.paragraphs)


@pytest.mark.parametrize("version", [f"3033-v{number}" for number in range(1, 7)])
@pytest.mark.parametrize("product_code", ["3033", "3149"])
def test_all_word_exports_use_updated_footers_without_changing_source(tmp_path, version, product_code):
    content = _document(version)
    content["sections"]["footnotes"].update(
        constituents="Constituent source.", analytics="Analytics source.",
    )
    before = deepcopy(content)
    report = _report()
    report.product_code = product_code
    destination = tmp_path / f"{product_code}-{version}.docx"

    render_docx(report, content, destination)

    assert content == before
    exported = Document(destination)
    for index, marker in ((0, "Source¹"), (2, "Constituent source."), (3, "Analytics source.")):
        paragraph = next(
            paragraph for paragraph in exported.sections[index].footer.paragraphs
            if marker in paragraph.text
        )
        for run in paragraph.runs:
            if run.text:
                assert run.font.name == "Calibri"
                assert run.font.size.pt == 8
                assert run.bold is False
    with ZipFile(destination) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
        assert "Source¹" not in body
        assert "Constituent source." not in body
        assert "Analytics source." not in body


def test_old_v4_word_moves_footnote_to_the_selected_review_page(tmp_path):
    content = _document("3033-v4")
    content["presentation"]["page_one"]["review_page_count"] = 2
    _element(content, "historical_performance").update(page=2, x=2, w=8)
    _element(content, "footnote:historical").update(page=2, x=2, w=8)
    before = deepcopy(content)
    destination = tmp_path / "old-v4-updated-footer.docx"

    render_docx(_report(), content, destination)

    assert content == before
    exported = Document(destination)
    assert all("Source¹" not in paragraph.text for paragraph in exported.sections[0].footer.paragraphs)
    paragraph = next(paragraph for paragraph in exported.sections[1].footer.paragraphs if "Source¹" in paragraph.text)
    assert paragraph.paragraph_format.left_indent.pt > 0
    assert paragraph.runs[0].font.name == "Calibri"
    assert paragraph.runs[0].font.size.pt == 8


def test_v6_real_calibri_pdf_and_responsive_ticker_cards(monkeypatch, tmp_path):
    import ctypes

    import pypdfium2 as pdfium
    from playwright.sync_api import sync_playwright
    from app.rendering.artifacts import _render_file, _verify_calibri_loaded

    font_path = settings.calibri_font_path or Path("C:/Windows/Fonts/calibri.ttf")
    if not font_path.is_file():
        pytest.skip("A licensed local Calibri Regular TTF is required for real-font verification.")
    monkeypatch.setattr(settings, "calibri_font_path", font_path)
    content = _document("3033-v6")
    content["sections"]["footnotes"].update(
        constituents="Constituent source.", analytics="Analytics source.",
    )
    html_path = tmp_path / "3033-v6.html"
    html_path.write_text(render_html(_report(), content), encoding="utf-8")

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        try:
            page = browser.new_page()
            page.route("https://**", lambda route: route.abort())
            for width in (1920, 1440, 1024, 390, 320):
                page.set_viewport_size({"width": width, "height": 1000})
                page.goto(html_path.as_uri())
                _verify_calibri_loaded(page, content)
                assert page.locator(".product-code").all_text_contents() == ["3033", "3149"]
                assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
                sidebar_box = page.locator(".product-sidebar").bounding_box()
                body_box = page.locator(".report-document").bounding_box()
                expected_position = "fixed" if width > 800 else "static"
                assert page.locator(".product-sidebar").evaluate("node => getComputedStyle(node).position") == expected_position
                for selector in (
                    ".reader-meta-label",
                    ".reader-meta-row",
                    ".reader-category",
                    ".share-heading",
                    ".related-products > h2",
                ):
                    assert page.locator(selector).first.evaluate(
                        "node => getComputedStyle(node).color"
                    ) == "rgb(0, 0, 0)"
                assert page.locator(".reader-author").evaluate(
                    "node => getComputedStyle(node).color"
                ) == "rgb(38, 96, 173)"
                if width > 800:
                    assert sidebar_box["x"] + sidebar_box["width"] <= body_box["x"]
                else:
                    assert sidebar_box["y"] + sidebar_box["height"] <= body_box["y"]
                page.locator(".product-link").first.focus()
                assert page.locator(".product-link").first.evaluate("node => node.matches(':focus-visible')")
                page.screenshot(path=str(tmp_path / f"html-{width}.png"))
                page.evaluate("window.scrollTo(0, document.documentElement.scrollHeight)")
                assert page.evaluate("window.scrollY") > 0
                scrolled_box = page.locator(".product-sidebar").bounding_box()
                if width > 800:
                    assert scrolled_box["x"] == pytest.approx(sidebar_box["x"], abs=0.1)
                    assert scrolled_box["y"] == pytest.approx(sidebar_box["y"], abs=0.1)
                else:
                    assert scrolled_box["y"] < sidebar_box["y"]
                page.screenshot(path=str(tmp_path / f"html-{width}-scrolled.png"))
                page.evaluate("window.scrollTo(0, 0)")
            page.emulate_media(media="print")
            assert not page.locator(".product-sidebar").is_visible()
            assert page.locator(".report-reader").evaluate("node => getComputedStyle(node).display") == "block"

            page.emulate_media(media="screen")
            page.set_viewport_size({"width": 1000, "height": 1200})
            for bottom in (16, 20, 36):
                _element(content, "footnote:historical")["footer_bottom_mm"] = bottom
                page.set_content(render_html(_report(), content, layout_mode="paged"))
                _verify_calibri_loaded(page, content)
                assert page.locator(".product-sidebar").count() == 0
                first_page = page.locator('.report-page[data-page="1"]')
                footnote = first_page.locator(".footnote")
                note_box = footnote.bounding_box()
                page_box = first_page.bounding_box()
                gap = (page_box["y"] + page_box["height"] - note_box["y"] - note_box["height"]) * 25.4 / 96
                assert gap == pytest.approx(bottom, abs=0.1)
                footer_box = first_page.locator(".page-footer").bounding_box()
                assert note_box["y"] + note_box["height"] <= footer_box["y"]
                if bottom == 20:
                    first_page.screenshot(path=str(tmp_path / "page-one.png"))
        finally:
            browser.close()

    pdf_path = tmp_path / "3033-v6.pdf"
    _render_file(_report(), SimpleNamespace(content=content), "pdf", pdf_path)
    exported = pdfium.PdfDocument(pdf_path)
    try:
        assert len(exported) == 5
        for page_index, text in ((0, "Source"), (2, "Constituent"), (3, "Analytics")):
            pdf_page = exported[page_index]
            text_page = pdf_page.get_textpage()
            search = text_page.search(text)
            try:
                match = search.get_next()
                assert match is not None
                flags = ctypes.c_int()
                size = pdfium.raw.FPDFText_GetFontInfo(text_page, match[0], None, 0, ctypes.byref(flags))
                buffer = ctypes.create_string_buffer(size)
                pdfium.raw.FPDFText_GetFontInfo(text_page, match[0], buffer, size, ctypes.byref(flags))
                assert "calibri" in buffer.value.decode().lower()
            finally:
                search.close()
                text_page.close()
                pdf_page.close()
    finally:
        exported.close()
    print(f"Visual evidence: {tmp_path}")
