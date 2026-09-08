from copy import deepcopy
from datetime import date

import pytest

from app.domain.document import initial_document
from app.domain.editorial_translation import ProtectedText, TranslationError, translate_review


def documents():
    source = initial_document("source", date(2026, 8, 31), "test-v1", "test-v1", "TEST", "TEST", "EN")
    target = initial_document("target", date(2026, 8, 31), "test-v1", "test-v1", "TEST", "TEST", "ZH_HANS")
    block = {"block_id": "summary", "type": "rich_text", "title": "Market review", "content": '<p>Return 12.5% <strong>HKD 100</strong></p>', "x": 0, "y": 0, "w": 12, "h": 4}
    source["sections"]["month_in_review"]["blocks"] = [block]
    target["sections"]["month_in_review"]["blocks"] = [{**block, "title": "", "content": ""}]
    return source, target


def convert(source, target, translator=None):
    return translate_review(source, target, source_id="source", target_id="target", source_version=1, source_language="EN", target_language="ZH_HANS", model="test", translate=translator or (lambda texts, *_: {key: value.replace("Return", "回报").replace("Market review", "市场回顾") for key, value in texts.items()}))


def test_translation_preserves_markup_numbers_and_manual_edits():
    source, target = documents()
    translated, preserved = convert(source, target)
    review = translated["sections"]["month_in_review"]
    assert review["title"] == "市场回顾"
    assert review["blocks"][0]["content"] == '<p>回报 12.5% <strong>HKD 100</strong></p>'
    assert not preserved
    unchanged, _ = convert(source, translated, lambda *_: pytest.fail("unchanged fields must not call provider"))
    assert unchanged == translated
    translated["sections"]["month_in_review"]["blocks"][0]["content"] = "<p>人工定稿文字</p>"
    source["sections"]["month_in_review"]["blocks"][0]["content"] = "<p>Changed return 10%</p>"
    revised, preserved = convert(source, translated, lambda *_: pytest.fail("manual text must not be translated"))
    assert revised["sections"]["month_in_review"]["blocks"][0]["content"] == "<p>人工定稿文字</p>"
    assert preserved == ["sections.month_in_review.blocks.summary.content"]


@pytest.mark.parametrize("response", ["Changed 99%", "[[KEEP_0]] [[KEEP_0]]", "<script>bad</script>", ""])
def test_translation_rejects_changed_facts_and_invalid_text(response):
    protected = ProtectedText("Return 12.5%", html=False)
    with pytest.raises(TranslationError):
        protected.restore({"0": response})


def test_only_review_text_is_sent():
    source, target = documents()
    source["sections"]["company_news"] = [{"summary": "paid news must remain private"}]
    source["sections"]["footnotes"] = {"historical": "approved disclosure"}
    before = deepcopy(source)

    def translator(texts, *_):
        assert "paid news" not in str(texts)
        assert "approved disclosure" not in str(texts)
        assert "12.5" not in str(texts)
        return {key: value.replace("Market review", "市场回顾").replace("Return", "回报") for key, value in texts.items()}

    convert(source, target, translator)
    assert source == before


def test_openai_adapter_uses_bounded_structured_request(monkeypatch):
    import httpx
    from app.core.config import settings
    from app.integrations.translation import translate_texts

    monkeypatch.setattr(settings, "translation_provider", "OPENAI_COMPATIBLE")
    monkeypatch.setattr(settings, "translation_base_url", "https://approved.example/v1")
    monkeypatch.setattr(settings, "translation_model", "approved-model")
    monkeypatch.setattr(settings, "translation_api_key", "test-only")
    real_client = httpx.Client

    def respond(request):
        import json
        payload = json.loads(request.content)
        assert payload["response_format"] == {"type": "json_object"}
        assert "tools" not in payload
        assert str(request.url) == "https://approved.example/v1/chat/completions"
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": '{"field": "市场"}'}}]})

    monkeypatch.setattr(httpx, "Client", lambda **kwargs: real_client(transport=httpx.MockTransport(respond), **kwargs))
    assert translate_texts({"field": "Market"}, "EN", "ZH_HANS") == {"field": "市场"}


def report_pair(client):
    source = client.post("/api/v1/reports", json={"report_date": "2026-08-31", "language_mode": "EN"}).json()
    detail = client.get(f"/api/v1/reports/{source['id']}").json()
    content = detail["latest_document"]["content"]
    content["sections"]["month_in_review"]["summary"] = "Market review"
    saved = client.patch(f"/api/v1/reports/{source['id']}/document", json={"version": 1, "content": content})
    assert saved.status_code == 200, saved.text
    target = client.post(f"/api/v1/reports/{source['id']}/language-variants", json={"source_document_version": 2, "language_mode": "ZH_HANS"}).json()
    return source, target


@pytest.fixture()
def enabled_translation(monkeypatch):
    from app.core.config import settings
    from app.domain.service import translations
    monkeypatch.setattr(settings, "translation_provider", "OPENAI_COMPATIBLE")
    monkeypatch.setattr(settings, "translation_base_url", "https://approved.example/v1")
    monkeypatch.setattr(settings, "translation_model", "approved-model")
    monkeypatch.setattr(settings, "translation_api_key", "test-only")
    monkeypatch.setattr(settings, "task_mode", "EAGER")
    monkeypatch.setattr(translations, "translate_texts", lambda texts, *_: {key: value.replace("Market review", "市场回顾") for key, value in texts.items()})


def test_translation_job_appends_once_and_can_be_replayed(client, enabled_translation):
    source, target = report_pair(client)
    url = f"/api/v1/reports/{source['id']}/language-variants/{target['id']}/translations"
    command = {"source_document_version": 2, "target_document_version": 1}
    response = client.post(url, json=command, headers={"Idempotency-Key": "one"})
    assert response.status_code == 202, response.text
    job = response.json()
    assert job["status"] == "SUCCEEDED", job
    current = client.get(f"/api/v1/reports/{target['id']}").json()
    assert current["latest_document"]["content"]["sections"]["month_in_review"]["summary"] == "市场回顾"
    assert current["latest_document"]["version"] == 2
    assert client.post(url, json=command, headers={"Idempotency-Key": "one"}).json()["id"] == job["id"]
    assert client.get(job["status_url"]).json()["status"] == "SUCCEEDED"
    assert client.get(f"/api/v1/reports/{target['id']}/translation-jobs/latest").json()["id"] == job["id"]


def test_translation_conflict_preserves_target(client, enabled_translation, monkeypatch):
    from app.domain.service import translations
    from app.api.routes import reports as routes
    from app.core.database import get_db

    source, target = report_pair(client)
    monkeypatch.setattr(routes, "dispatch_translation", lambda *_: None)
    job = client.post(f"/api/v1/reports/{source['id']}/language-variants/{target['id']}/translations", json={"source_document_version": 2, "target_document_version": 1}, headers={"Idempotency-Key": "queued"}).json()
    detail = client.get(f"/api/v1/reports/{target['id']}").json()
    content = detail["latest_document"]["content"]
    content["sections"]["month_in_review"]["summary"] = "人工修改"
    assert client.patch(f"/api/v1/reports/{target['id']}/document", json={"version": 1, "content": content}).status_code == 200
    with next(client.app.dependency_overrides[get_db]()) as db:
        assert translations.execute_translation(db, job["id"]) == "FAILED"
    assert client.get(job["status_url"]).json()["error"]["error_code"] == "VERSION_CONFLICT"
    assert client.get(f"/api/v1/reports/{target['id']}").json()["latest_document"]["version"] == 2


def test_custom_title_replaces_only_target_template_title():
    source, target = documents()
    target["sections"]["month_in_review"]["blocks"][0]["title"] = "月度回顾"
    result, _ = convert(source, target)
    assert result["sections"]["month_in_review"]["display_title"] == "市场回顾"


def test_legacy_source_updates_canonical_target_blocks():
    source, target = documents()
    del source["sections"]["month_in_review"]["blocks"]
    source["sections"]["month_in_review"]["summary"] = "Market review"
    target["sections"]["month_in_review"]["blocks"][0]["title"] = "月度回顾"
    result, _ = convert(source, target)
    assert result["sections"]["month_in_review"]["blocks"][0]["content"] == "<p>市场回顾</p>"
    assert result["sections"]["month_in_review"]["summary"] == "市场回顾"


def test_sync_preserves_target_only_blocks():
    from app.domain.service.reports import _copy_review_layout
    source, target = documents()
    target["sections"]["month_in_review"]["blocks"].append({"block_id": "manual", "type": "rich_text", "title": "人工标题", "content": "<p>人工文字</p>", "x": 0, "y": 4, "w": 12, "h": 4})
    _copy_review_layout(source, target, "ZH_HANS")
    result, preserved = convert(source, target)
    assert result["sections"]["month_in_review"]["blocks"][1]["content"] == "<p>人工文字</p>"
    assert "blocks.manual" in preserved


def test_provider_failure_never_appends_or_exposes_body(client, enabled_translation, monkeypatch):
    from app.domain.service import translations
    source, target = report_pair(client)
    monkeypatch.setattr(translations, "translate_texts", lambda *_: (_ for _ in ()).throw(RuntimeError("private-provider-body")))
    response = client.post(f"/api/v1/reports/{source['id']}/language-variants/{target['id']}/translations", json={"source_document_version": 2, "target_document_version": 1}, headers={"Idempotency-Key": "failure"})
    assert response.json()["status"] == "FAILED"
    assert "private-provider-body" not in response.text
    assert client.get(f"/api/v1/reports/{target['id']}").json()["latest_document"]["version"] == 1


def test_unknown_numbers_fail_before_applying(client, enabled_translation):
    source, target = report_pair(client)
    current = client.get(f"/api/v1/reports/{source['id']}").json()
    content = current["latest_document"]["content"]
    content["sections"]["month_in_review"]["summary"] = "Market review 99%"
    assert client.patch(f"/api/v1/reports/{source['id']}/document", json={"version": 2, "content": content}).status_code == 200
    response = client.post(f"/api/v1/reports/{source['id']}/language-variants/{target['id']}/translations", json={"source_document_version": 3, "target_document_version": 1}, headers={"Idempotency-Key": "numbers"})
    assert response.json()["error"]["error_code"] == "TRANSLATION_UNBOUND_NUMBER"
    assert client.get(f"/api/v1/reports/{target['id']}").json()["latest_document"]["version"] == 1


def test_wrong_language_is_not_applied():
    source, target = documents()
    with pytest.raises(TranslationError, match="TRANSLATION_LANGUAGE_MISMATCH"):
        convert(source, target, lambda texts, *_: texts)


@pytest.mark.parametrize("source_language,target_language,original,expected", [("EN", "ZH_HANT", "Market review", "市場回顧"), ("ZH_HANS", "EN", "市场回顾", "Market review"), ("ZH_HANT", "EN", "市場回顧", "Market review")])
def test_translation_supports_each_english_chinese_direction(source_language, target_language, original, expected):
    source = initial_document("source", date(2026, 8, 31), "test-v1", "test-v1", "TEST", "TEST", source_language)
    target = initial_document("target", date(2026, 8, 31), "test-v1", "test-v1", "TEST", "TEST", target_language)
    source["sections"]["month_in_review"]["summary"] = original
    translated, _ = translate_review(source, target, source_id="source", target_id="target", source_version=1, source_language=source_language, target_language=target_language, model="test", translate=lambda texts, *_: {key: value.replace(original, expected) for key, value in texts.items()})
    assert translated["sections"]["month_in_review"]["summary"] == expected


def test_translated_document_is_shared_by_html_pdf_and_docx(client, enabled_translation):
    from io import BytesIO
    import pypdfium2 as pdfium
    from docx import Document
    from sqlalchemy import select
    from app.core.database import get_db
    from app.domain.models import RenderArtifact

    source, target = report_pair(client)
    translated = client.post(f"/api/v1/reports/{source['id']}/language-variants/{target['id']}/translations", json={"source_document_version": 2, "target_document_version": 1}, headers={"Idempotency-Key": "outputs"}).json()
    assert translated["status"] == "SUCCEEDED", translated
    document = client.get(f"/api/v1/reports/{target['id']}").json()["latest_document"]
    assert client.post(f"/api/v1/reports/{target['id']}/finalize", json={"version": document["version"]}).status_code == 200
    jobs = client.post(f"/api/v1/reports/{target['id']}/renders", json={"formats": ["html", "pdf", "docx"]}, headers={"Idempotency-Key": "translated-output"}).json()
    for job in jobs:
        assert job["status"] == "SUCCEEDED", job
        signed = client.get(f"/api/v1/artifacts/{job['artifact_id']}/download").json()
        body = client.get(signed["download_url"]).content
        if job["format"] == "html":
            assert "市场回顾" in body.decode("utf-8")
        elif job["format"] == "docx":
            output = Document(BytesIO(body))
            assert "市场回顾" in " ".join(paragraph.text for paragraph in output.paragraphs)
        else:
            output = pdfium.PdfDocument(body)
            try:
                assert "市场回顾" in " ".join(page.get_textpage().get_text_bounded() for page in output)
            finally:
                output.close()
    with next(client.app.dependency_overrides[get_db]()) as db:
        artifacts = list(db.scalars(select(RenderArtifact).where(RenderArtifact.report_id == target["id"])))
        assert len(artifacts) == 3
        assert all(artifact.document_version == document["version"] and artifact.content_manifest["language_mode"] == "ZH_HANS" and artifact.content_manifest["document_checksum"] == document["checksum"] for artifact in artifacts)